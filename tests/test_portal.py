"""Client document requests and the client portal."""

import io
import re

import pytest
from openpyxl import load_workbook

from app import db
from app.main import app
from tests.test_auth import add_user, admin, at_once, login

pytestmark = pytest.mark.usefixtures("real_auth")

PDF = ("bank.pdf", b"%PDF-1.4 statement", "application/pdf")


def add_client_user(a, client_id, email="client@example.com", name="Cara Client"):
    r = a.post("/api/users", json={"email": email, "name": name, "role": "client",
                                   "client_id": client_id, "password": "long enough pw"})
    assert r.status_code == 201, r.text
    return login(email, "long enough pw")


def setup():
    """Admin, preparer, reviewer, a client user of Acme, and an Acme engagement."""
    a = admin()
    prep, rev = add_user(a, "preparer"), add_user(a, "reviewer")
    acme = a.post("/api/clients", json={"name": "Acme"}).json()
    eid = a.post(f"/api/clients/{acme['id']}/engagements",
                 json={"name": "FY2026", "period_end": "2026-12-31"}).json()["id"]
    cl = add_client_user(a, acme["id"])
    return a, prep, rev, cl, acme["id"], eid


def ask(c, eid, title="Bank statements for December", **extra):
    r = c.post(f"/api/engagements/{eid}/requests", json={"title": title, **extra})
    assert r.status_code == 201, r.text
    return r.json()


def upload(c, rid, file=PDF):
    return c.post(f"/api/requests/{rid}/files", files={"file": file})


# Endpoints a client user may call; everything else must refuse them.
CLIENT_PATHS = {"/", "/api/auth/status", "/api/auth/setup", "/api/auth/login", "/api/auth/logout",
                "/api/auth/me", "/api/auth/password", "/api/health", "/api/portal/requests",
                "/api/requests/{request_id}/files", "/api/requests/{request_id}/messages",
                "/api/request-files/{file_id}"}


def test_client_users_are_refused_everywhere_but_the_portal():
    _, _, _, cl, _, eid = setup()
    checked = 0
    for path, methods in app.openapi()["paths"].items():
        if path in CLIENT_PATHS:
            continue
        url = re.sub(r"{[^}]+}", str(eid), path)
        for method in methods:
            r = cl.request(method.upper(), url)
            assert r.status_code == 403, (method, path, r.status_code)
            checked += 1
    assert checked > 50


def test_client_role_needs_a_client_and_staff_roles_cannot_have_one():
    a = admin()
    acme = a.post("/api/clients", json={"name": "Acme"}).json()
    body = {"email": "x@example.com", "name": "X", "password": "long enough pw"}
    r = a.post("/api/users", json={**body, "role": "client"})
    assert r.status_code == 400 and "client" in r.json()["detail"]
    r = a.post("/api/users", json={**body, "role": "client", "client_id": 999})
    assert r.status_code == 400
    r = a.post("/api/users", json={**body, "role": "preparer", "client_id": acme["id"]})
    assert r.status_code == 400
    r = a.post("/api/users", json={**body, "role": "client", "client_id": acme["id"]})
    assert r.status_code == 201
    u = r.json()
    assert u["client_id"] == acme["id"] and u["client_name"] == "Acme" and u["rank"] == 0
    log = a.get("/api/activity", params={"action": "user.create"}).json()
    assert log[0]["summary"].endswith("as Client of Acme")
    # Becoming a preparer drops the link; becoming a client again needs a client.
    r = a.patch(f"/api/users/{u['id']}", json={"role": "preparer"})
    assert r.status_code == 200 and r.json()["client_id"] is None
    assert a.patch(f"/api/users/{u['id']}", json={"role": "client"}).status_code == 400
    r = a.patch(f"/api/users/{u['id']}", json={"role": "client", "client_id": acme["id"]})
    assert r.json()["client_name"] == "Acme"


def test_client_can_sign_in_see_themselves_and_change_password():
    _, _, _, cl, _, _ = setup()
    me = cl.get("/api/auth/me").json()
    assert me["role"] == "client" and me["client_name"] == "Acme"
    r = cl.post("/api/auth/password", json={"current_password": "long enough pw",
                                            "new_password": "another long pw"})
    assert r.status_code == 204
    assert cl.get("/api/activity", params={"engagement_id": 1}).status_code == 403


def test_request_flow_from_ask_to_accept():
    a, prep, rev, cl, _, eid = setup()
    req = ask(prep, eid, details="All bank accounts", due_date="2026-01-15")
    assert req["number"] == 1 and req["status"] == "open" and req["overdue"]

    portal = cl.get("/api/portal/requests").json()
    assert portal["client"] == "Acme"
    [eng] = portal["engagements"]
    assert eng["name"] == "FY2026" and eng["requests"][0]["title"] == req["title"]

    r = upload(cl, req["id"])
    assert r.status_code == 201
    req = r.json()
    assert req["status"] == "submitted" and not req["overdue"]
    [f] = req["files"]
    assert f["from_client"] and f["uploaded_by"] == "Cara Client"

    r = prep.get(f"/api/request-files/{f['id']}")
    assert r.content == PDF[1] and r.headers["content-type"] == "application/octet-stream"
    assert "attachment" in r.headers["content-disposition"]

    r = prep.post(f"/api/requests/{req['id']}/status", json={"status": "returned"})
    assert r.status_code == 400  # needs a message
    r = prep.post(f"/api/requests/{req['id']}/status",
                  json={"status": "returned", "message": "November is missing"})
    assert r.json()["status"] == "returned"
    assert cl.get("/api/portal/requests").json()["engagements"][0]["requests"][0][
        "messages"][-1]["text"] == "November is missing"

    upload(cl, req["id"], ("nov.pdf", b"%PDF nov", "application/pdf"))
    r = prep.post(f"/api/requests/{req['id']}/status", json={"status": "accepted"})
    assert r.status_code == 200 and r.json()["status"] == "accepted"
    assert [m["status_to"] for m in r.json()["messages"]] == [
        "submitted", "returned", "submitted", "accepted"]

    # Once accepted, the client can't add or remove files.
    r = upload(cl, req["id"])
    assert r.status_code == 400 and "accepted" in r.json()["detail"]
    assert cl.delete(f"/api/request-files/{f['id']}").status_code == 400

    actions = [x["action"] for x in a.get("/api/activity",
                                           params={"engagement_id": eid}).json()]
    for action in ("request.create", "request.upload", "request.file_download",
                   "request.return", "request.accept"):
        assert action in actions
    uploads = a.get("/api/activity", params={"engagement_id": eid,
                                             "action": "request.upload"}).json()
    assert uploads[-1]["user_role"] == "client" and uploads[-1]["details"]["sha256"]


def test_accept_needs_a_submitted_request():
    _, prep, _, _, _, eid = setup()
    req = ask(prep, eid)
    r = prep.post(f"/api/requests/{req['id']}/status", json={"status": "accepted"})
    assert r.status_code == 400
    r = prep.post(f"/api/requests/{req['id']}/status", json={"status": "returned",
                                                             "message": "x"})
    assert r.status_code == 400


def test_clients_only_see_their_own_clients_requests():
    a, prep, _, cl, _, eid = setup()
    other = a.post("/api/clients", json={"name": "Globex"}).json()
    oeid = a.post(f"/api/clients/{other['id']}/engagements",
                  json={"name": "FY2026", "period_end": "2026-12-31"}).json()["id"]
    theirs = ask(prep, oeid, "Globex payroll")
    upload(prep, theirs["id"])
    fid = prep.get(f"/api/engagements/{oeid}/requests").json()["requests"][0]["files"][0]["id"]
    ask(prep, eid, "Acme fixed assets")

    portal = cl.get("/api/portal/requests").json()
    titles = [r["title"] for e in portal["engagements"] for r in e["requests"]]
    assert titles == ["Acme fixed assets"]
    assert upload(cl, theirs["id"]).status_code == 404
    assert cl.get(f"/api/request-files/{fid}").status_code == 404
    assert cl.delete(f"/api/request-files/{fid}").status_code == 404
    assert cl.post(f"/api/requests/{theirs['id']}/messages",
                   json={"text": "hi"}).status_code == 404
    assert cl.get(f"/api/engagements/{oeid}/requests").status_code == 403


def test_who_can_remove_files_and_delete_requests():
    _, prep, rev, cl, _, eid = setup()
    req = ask(prep, eid)
    upload(cl, req["id"])
    fid = upload(prep, req["id"], ("mine.pdf", b"from the team", "application/pdf")).json()[
        "files"][1]["id"]
    client_fid = prep.get(f"/api/engagements/{eid}/requests").json()["requests"][0][
        "files"][0]["id"]

    assert cl.delete(f"/api/request-files/{fid}").status_code == 403  # not theirs
    assert prep.delete(f"/api/request-files/{client_fid}").status_code == 403  # needs review
    assert prep.delete(f"/api/request-files/{fid}").status_code == 204  # own file
    assert cl.delete(f"/api/request-files/{client_fid}").status_code == 204
    r = prep.get(f"/api/engagements/{eid}/requests").json()["requests"][0]
    assert r["files"] == [] and r["status"] == "open"  # nothing left to look at

    upload(cl, req["id"])
    assert prep.delete(f"/api/requests/{req['id']}").status_code == 403  # has client files
    rev.post(f"/api/requests/{req['id']}/status", json={"status": "accepted"})
    r = rev.delete(f"/api/requests/{req['id']}")  # accepted evidence is locked
    assert r.status_code == 400 and "Return it" in r.json()["detail"]
    rev.post(f"/api/requests/{req['id']}/status",
             json={"status": "returned", "message": "Wrong year"})
    assert rev.delete(f"/api/requests/{req['id']}").status_code == 204
    empty = ask(prep, eid, "Something else")
    assert prep.delete(f"/api/requests/{empty['id']}").status_code == 204


def test_request_numbers_are_never_reused():
    _, prep, _, _, _, eid = setup()
    ask(prep, eid, "One")
    two = ask(prep, eid, "Two")
    prep.delete(f"/api/requests/{two['id']}")
    assert ask(prep, eid, "Three")["number"] == 3


def test_concurrent_requests_get_different_numbers():
    _, prep, _, _, _, eid = setup()
    results = at_once(4, lambda i: prep.post(f"/api/engagements/{eid}/requests",
                                             json={"title": f"Doc {i}"}))
    assert all(r.status_code == 201 for r in results), [r.text for r in results]
    assert sorted(r.json()["number"] for r in results) == [1, 2, 3, 4]


def test_messages_both_ways_and_edits():
    _, prep, _, cl, _, eid = setup()
    req = ask(prep, eid)
    assert cl.post(f"/api/requests/{req['id']}/messages",
                   json={"text": "Which accounts?"}).status_code == 201
    r = prep.post(f"/api/requests/{req['id']}/messages", json={"text": "All of them"})
    msgs = r.json()["messages"]
    assert [(m["author"], m["from_client"]) for m in msgs] == [("Cara Client", True),
                                                               ("Preparer", False)]
    r = prep.patch(f"/api/requests/{req['id']}", json={"title": "All bank statements",
                                                      "due_date": None})
    assert r.json()["title"] == "All bank statements"
    assert prep.patch(f"/api/requests/{req['id']}", json={"title": None}).status_code == 400
    assert cl.patch(f"/api/requests/{req['id']}", json={"title": "x"}).status_code == 403


def test_export_escapes_formulas():
    _, prep, _, cl, _, eid = setup()
    req = ask(prep, eid, "=HYPERLINK(\"http://evil\")", details="+cmd|' /C calc'!A0")
    upload(cl, req["id"], ("=evil.pdf", b"x", "application/pdf"))
    r = prep.get(f"/api/engagements/{eid}/requests/export")
    assert r.status_code == 200 and "Acme-FY2026-requests.xlsx" in r.headers[
        "content-disposition"]
    book = load_workbook(io.BytesIO(r.content))
    ws = book["Requests"]
    assert ws["B2"].value.startswith("=HYPERLINK") and ws["B2"].data_type == "s"
    assert book["Files"]["B2"].data_type == "s"
    r = prep.get(f"/api/engagements/{eid}/requests/export", params={"format": "csv"})
    assert "'=HYPERLINK" in r.text and "'+cmd" in r.text and "'=evil.pdf" in r.text
    assert prep.get(f"/api/engagements/{eid}/requests/export",
                    params={"format": "pdf"}).status_code == 400


def test_deleting_the_client_unlinks_its_portal_users_and_requests():
    a, prep, _, cl, cid, eid = setup()
    req = ask(prep, eid)
    upload(cl, req["id"])
    partner = add_user(a, "partner")
    assert partner.delete(f"/api/clients/{cid}").status_code == 204
    assert cl.get("/api/portal/requests").json() == {"client": None, "engagements": []}
    with db.SessionLocal() as s:
        assert s.query(db.DocumentRequest).count() == 0
        assert s.query(db.RequestFile).count() == 0
        assert s.query(db.ClientUser).count() == 0


def test_team_list_shows_who_at_the_client_has_access():
    a, prep, _, _, _, eid = setup()
    data = prep.get(f"/api/engagements/{eid}/requests").json()
    assert data["contacts"] == [{"name": "Cara Client", "email": "client@example.com",
                                 "active": True}]
    assert data["summary"]["total"] == 0


# --- Review fixes ---------------------------------------------------------------


def accepted_request(prep, rev, cl, eid):
    req = ask(prep, eid)
    upload(cl, req["id"])
    r = rev.post(f"/api/requests/{req['id']}/status", json={"status": "accepted"})
    assert r.status_code == 200
    return r.json()


def test_accepted_requests_are_locked_until_a_reviewer_returns_them():
    _, prep, rev, cl, _, eid = setup()
    req = accepted_request(prep, rev, cl, eid)
    r = prep.patch(f"/api/requests/{req['id']}", json={"title": "Signed board minutes"})
    assert r.status_code == 400 and "Return" in r.json()["detail"]
    back = {"status": "returned", "message": "Wrong year"}
    assert prep.post(f"/api/requests/{req['id']}/status", json=back).status_code == 403
    assert rev.post(f"/api/requests/{req['id']}/status", json=back).json()["status"] == "returned"
    # Returned by mistake: it can be accepted again without a new upload.
    r = rev.post(f"/api/requests/{req['id']}/status", json={"status": "accepted"})
    assert r.json()["status"] == "accepted"


def test_moving_a_client_user_is_logged_and_signs_out_role_switches():
    a, _, _, cl, cid, _ = setup()
    globex = a.post("/api/clients", json={"name": "Globex"}).json()
    uid = cl.get("/api/auth/me").json()["id"]
    r = a.patch(f"/api/users/{uid}", json={"client_id": globex["id"]})
    assert r.json()["client_name"] == "Globex"
    log = a.get("/api/activity", params={"action": "user.update"}).json()[0]
    assert "client to Globex" in log["summary"] and log["details"]["client_id"] == globex["id"]
    assert cl.get("/api/auth/me").status_code == 200  # moving clients keeps the session
    a.patch(f"/api/users/{uid}", json={"role": "preparer"})
    assert cl.get("/api/auth/me").status_code == 401  # becoming staff means signing in again


def test_file_ownership_stays_on_the_senders_side():
    a, prep, rev, cl, cid, eid = setup()
    req = ask(prep, eid)
    team_file = upload(prep, req["id"], ("team.pdf", b"team", "application/pdf")).json()[
        "files"][0]["id"]
    # The preparer becomes a client of Acme: the file they sent as staff isn't theirs there.
    pid = prep.get("/api/auth/me").json()["id"]
    a.patch(f"/api/users/{pid}", json={"role": "client", "client_id": cid})
    turned = login("preparer@example.com", "long enough pw")
    assert turned.delete(f"/api/request-files/{team_file}").status_code == 403
    # And a client who becomes a preparer can't remove what they sent as the client.
    client_file = upload(cl, req["id"]).json()["files"][-1]["id"]
    uid = cl.get("/api/auth/me").json()["id"]
    a.patch(f"/api/users/{uid}", json={"role": "preparer"})
    staff = login("client@example.com", "long enough pw")
    assert staff.delete(f"/api/request-files/{client_file}").status_code == 403
    assert rev.delete(f"/api/request-files/{client_file}").status_code == 204


def test_other_clients_ids_look_missing_and_deactivated_clients_are_out():
    a, prep, _, cl, _, eid = setup()
    other = a.post("/api/clients", json={"name": "Globex"}).json()
    oeid = a.post(f"/api/clients/{other['id']}/engagements",
                  json={"name": "FY2026", "period_end": "2026-12-31"}).json()["id"]
    theirs = ask(prep, oeid)
    fid = upload(prep, theirs["id"]).json()["files"][0]["id"]
    for path, missing in ((f"/api/request-files/{fid}", f"/api/request-files/{fid + 99}"),):
        for method in ("GET", "DELETE"):
            assert cl.request(method, path).json() == cl.request(method, missing).json()
    for method, sub in (("POST", "messages"), ("POST", "files")):
        kw = {"json": {"text": "x"}} if sub == "messages" else {"files": {"file": PDF}}
        r1 = cl.request(method, f"/api/requests/{theirs['id']}/{sub}", **kw)
        r2 = cl.request(method, f"/api/requests/{theirs['id'] + 99}/{sub}", **kw)
        assert r1.status_code == r2.status_code == 404 and r1.json() == r2.json()
    uid = cl.get("/api/auth/me").json()["id"]
    a.patch(f"/api/users/{uid}", json={"active": False})
    assert cl.get("/api/portal/requests").status_code == 401


def test_uploads_over_the_limit_are_refused(monkeypatch):
    from app import portal
    _, prep, _, cl, _, eid = setup()
    monkeypatch.setattr(portal, "MAX_FILE_BYTES", 10)
    req = ask(prep, eid)
    assert upload(cl, req["id"], ("big.pdf", b"x" * 11, "application/pdf")).status_code == 413
    assert upload(cl, req["id"], ("ok.pdf", b"x" * 10, "application/pdf")).status_code == 201


def test_deleting_an_engagement_logs_the_client_documents():
    a, prep, _, cl, _, eid = setup()
    upload(cl, ask(prep, eid)["id"])
    partner = add_user(a, "partner")
    assert partner.delete(f"/api/engagements/{eid}").status_code == 204
    log = a.get("/api/activity", params={"action": "engagement.delete"}).json()[0]
    assert log["details"]["request_files"][0]["filename"] == "bank.pdf"


def test_oversized_bodies_are_refused_as_they_arrive(monkeypatch):
    from app import portal
    _, prep, _, cl, _, eid = setup()
    monkeypatch.setattr(portal, "MAX_FILE_BYTES", 1000)
    req = ask(prep, eid)
    big = b"x" * 200_000  # past the limit plus the room for form fields
    r = upload(cl, req["id"], ("big.pdf", big, "application/pdf"))
    assert r.status_code == 413 and "too large" in r.json()["detail"]

    def chunks():  # no Content-Length: counted as it streams
        yield b"--b\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a\"\r\n\r\n"
        for _ in range(20):
            yield b"x" * 10_000
        yield b"\r\n--b--\r\n"

    r = cl.post(f"/api/requests/{req['id']}/files", content=chunks(),
                headers={"Content-Type": "multipart/form-data; boundary=b"})
    assert r.status_code == 413 and "too large" in r.json()["detail"]
    # Other endpoints keep the general limit.
    assert prep.post(f"/api/requests/{req['id']}/messages",
                     json={"text": "x" * 2000}).status_code == 201
