"""Bank reconciliation, and rolling an engagement forward to next year."""

import csv
import io
import re
from datetime import date

import openpyxl
import pandas as pd
import pytest

from app import bankrec, dashboard, loader, rollforward
from tests.test_auth import add_user, admin, join, setup_engagement
from tests.test_statements import DATA, client, engagement, map_all, post_entry

PE = date(2026, 12, 31)
BANK = DATA / "sample_bank_2026.csv"


def frame(rows, cols):
    return pd.DataFrame([dict(zip(cols, r, strict=True)) for r in rows])


def bank_lines(rows):
    """Bank lines as date, description, money in (positive) or out."""
    df = frame([(i + 2, pd.Timestamp(d), t, a) for i, (d, t, a) in enumerate(rows)],
               ["row", "date", "description", "amount"])
    return bankrec.bank_frame(df, signed_amount=True, flip=False)


def ledger_lines(rows):
    df = frame([(i + 2, pd.Timestamp(d), f"JE{i}", "1000 Cash", t, a)
                for i, (d, t, a) in enumerate(rows)],
               ["row", "date", "entry_id", "account", "description", "amount"])
    return bankrec.ledger_frame(df, ["1000 Cash"])


def reconcile(bank, ledger, **kw):
    return bankrec.reconcile(bank_lines(bank), ledger_lines(ledger),
                             bankrec.Settings(PE, ["1000 Cash"], **kw))


def test_matches_by_amount_within_the_date_window():
    r = reconcile(
        bank=[("2026-03-03", "TRANSFER IN", 100.0), ("2026-03-20", "CHQ 1", -40.0),
              ("2026-03-21", "CHQ 2", -40.0), ("2026-04-30", "CHARGES", -5.0)],
        ledger=[("2026-03-01", "Receipt", 100.0), ("2026-03-19", "Pay A", -40.0),
                ("2026-03-21", "Pay B", -40.0), ("2026-02-01", "Old receipt", 100.0)],
        window_days=3, bank_balance=55.0, book_balance=160.0)
    s = r["summary"]
    assert s["matched"] == 3 and s["outstanding"] == 1 and s["bank_only"] == 1
    # The same amount twice: each ledger line takes the bank line nearest its date.
    days = {m["ledger_description"]: m["days"] for m in r["matched"]}
    assert days == {"Receipt": 2, "Pay A": 1, "Pay B": 0}
    # The February receipt is a month off any bank line of that amount: still outstanding.
    o = r["outstanding"][0]
    assert o["kind"] == "Deposit in transit" and o["description"] == "Old receipt"
    assert o["days_outstanding"] == (PE - date(2026, 2, 1)).days and o["stale"]
    st = r["statement"]
    assert st["deposits_in_transit"] == 100.0 and st["payments_not_in_books"] == 5.0
    assert st["adjusted_bank"] == 155.0 and st["adjusted_book"] == 155.0
    assert st["difference"] == 0.0


def test_items_cleared_after_the_period_end():
    r = reconcile(
        bank=[("2026-12-10", "CHQ 1", -10.0), ("2027-01-04", "TRANSFER IN", 250.0),
              ("2027-01-06", "CHQ 2", -75.5), ("2027-03-20", "CHQ 3", -9.0),
              ("2027-01-10", "LATER", 1.0)],
        ledger=[("2026-12-10", "Pay", -10.0), ("2026-12-30", "Late receipt", 250.0),
                ("2026-12-29", "Cheque", -75.5), ("2026-12-28", "Slow cheque", -9.0),
                ("2027-01-02", "Next year", 33.0), ("2026-01-01", "Opening balance", 900.0)])
    s = r["summary"]
    assert s["matched"] == 1 and s["outstanding"] == 3 and s["cleared_after"] == 2
    assert s["ledger_after_period"] == 1 and s["opening_lines_skipped"] == 1
    assert s["bank_after_period_unmatched"] == 2 and s["bank_only"] == 0
    cleared = {o["description"]: o["cleared"] for o in r["outstanding"]}
    assert cleared["Late receipt"]["date"] == "2027-01-04"
    assert cleared["Cheque"]["description"] == "CHQ 2"
    assert cleared["Slow cheque"] is None  # more than 60 days after the period end
    # A matched line isn't used twice, and nothing in the period clears an outstanding item.
    assert r["statement"]["difference"] is None  # no balances given
    assert r["statement"]["outstanding_payments"] == 84.5


def test_reconcile_errors():
    with pytest.raises(bankrec.BankRecError, match="cash account"):
        bankrec.reconcile(bank_lines([("2026-01-01", "x", 1.0)]),
                          ledger_lines([("2026-01-01", "x", 1.0)]), bankrec.Settings(PE, []))
    with pytest.raises(bankrec.BankRecError, match="up to the period end"):
        reconcile([("2026-01-01", "x", 1.0)], [("2027-01-01", "x", 1.0)])
    with pytest.raises(bankrec.BankRecError, match="no dated lines"):
        reconcile([("2026-01-01", "x", 0.0)], [("2026-01-01", "x", 1.0)])
    with pytest.raises(bankrec.BankRecError, match="between 0 and 60"):
        reconcile([("2026-01-01", "x", 1.0)], [("2026-01-01", "x", 1.0)], window_days=61)


def test_bank_statement_columns_and_sign():
    # Withdrawals and deposits are a debit and a credit on the bank's side; the running
    # balance beside them is not the amount.
    loaded = loader.load("bank.csv", BANK.read_bytes())
    assert loaded.report.mapping == {"date": "Date", "description": "Description",
                                     "debit": "Withdrawals", "credit": "Deposits"}
    b = bankrec.bank_frame(loaded.df, signed_amount=False, flip=False)
    first = b[b["description"].str.startswith("TRANSFER IN")].iloc[0]
    assert first["amount"] > 0
    # A single signed Amount column is money in positive, unless the user says otherwise.
    signed = loader.load("s.csv", b"Date,Details,Amount\n2026-01-02,FEE,-5\n")
    assert bankrec.bank_frame(signed.df, True, False)["amount"].tolist() == [-5.0]
    assert bankrec.bank_frame(signed.df, True, True)["amount"].tolist() == [5.0]
    assert loader.suggest_mapping(["Date", "Amount", "Balance"])["amount"] == "Amount"
    assert loader.suggest_mapping(["Date", "Balance"])["amount"] == "Balance"


def test_cash_accounts_are_suggested():
    accounts = ["1000 Cash", "1010 Cash at Bank - CIB", "1200 Receivables", "6400 Travel"]
    assert bankrec.cash_accounts(accounts, set()) == ["1000 Cash", "1010 Cash at Bank - CIB"]
    assert bankrec.cash_accounts(accounts, {"6400"}) == ["1000 Cash", "1010 Cash at Bank - CIB",
                                                         "6400 Travel"]


def closing_balance():
    rows = list(csv.DictReader(BANK.read_text().splitlines()))
    return float([r for r in rows if r["Date"] <= "2026-12-31"][-1]["Balance"])


def upload(eid, path, kind):
    r = client.post(f"/api/engagements/{eid}/uploads", data={"kind": kind},
                    files={"file": (path.name, path.read_bytes(), "text/csv")})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_reconcile_the_sample_statement():
    eid = engagement(with_prior=False)  # this year's trial balance: cash 389,150.00
    gl = upload(eid, DATA / "sample_gl.csv", "gl")
    bank = upload(eid, BANK, "bank")
    steps = lambda: [s["text"] for s in client.get(  # noqa: E731
        f"/api/engagements/{eid}/dashboard").json()["next_steps"]]
    assert "Reconcile the bank statement with the general ledger." in steps()

    o = client.get(f"/api/engagements/{eid}/bank-reconciliations/options").json()
    assert [u["id"] for u in o["bank_uploads"]] == [bank] and o["gl_upload_id"] == gl
    assert [a["account"] for a in o["accounts"] if a["cash"]] == ["1000 Cash"]
    assert o["tb_balances"]["1000"] == 389150.0

    body = {"bank_upload_id": bank, "gl_upload_id": gl, "accounts": ["1000 Cash"],
            "bank_balance": closing_balance()}
    r = client.post(f"/api/engagements/{eid}/bank-reconciliations", json=body)
    assert r.status_code == 201, r.text
    rec = r.json()
    res = rec["result"]
    assert res["settings"]["book_balance"] == 389150.0
    assert res["settings"]["book_balance_source"] == "trial balance"
    assert res["statement"]["difference"] == 0.0
    s = res["summary"]
    assert s["matched"] + s["outstanding"] == s["ledger_lines"] == 290
    assert s["stale"] == 1 and s["cleared_after"] == s["outstanding"] - 1
    assert {b["description"] for b in res["bank_only"]} == {"BANK CHARGES", "INTEREST CREDIT"}
    assert "Reconcile the bank statement with the general ledger." not in steps()
    assert "Follow up 1 old reconciling item on the bank reconciliation." in steps()

    # A difference is a high-priority step.
    r = client.post(f"/api/engagements/{eid}/bank-reconciliations",
                    json={**body, "book_balance": 1000.0})
    assert r.json()["difference"] != 0
    assert "The bank reconciliation has an unexplained difference; look into it." in steps()

    x = client.get(f"/api/bank-reconciliations/{rec['id']}/report")
    assert x.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Reconciliation", "Outstanding items", "Bank only", "Matched"]
    items = dict(wb["Reconciliation"].iter_rows(min_row=2, values_only=True))
    assert items["Unexplained difference"] == 0 and items["Client"] == "Acme"
    log = client.get(f"/api/activity?engagement_id={eid}&action=bankrec").json()
    assert [a["action"] for a in log] == ["bankrec.download", "bankrec.create", "bankrec.create"]
    assert log[-1]["summary"].startswith("Reconciled sample_bank_2026.csv with 1000 Cash: ")


def test_bank_uploads_and_their_reconciliations():
    eid = engagement(with_prior=False)
    gl = upload(eid, DATA / "sample_gl.csv", "gl")
    bank = upload(eid, BANK, "bank")
    tb = client.get(f"/api/engagements/{eid}").json()["uploads"][0]["id"]
    r = client.post(f"/api/engagements/{eid}/runs", json={"upload_id": bank})
    assert r.status_code == 400 and "bank reconciliation panel" in r.json()["detail"]
    for wrong in ({"bank_upload_id": gl, "gl_upload_id": gl},
                  {"bank_upload_id": bank, "gl_upload_id": tb}):
        r = client.post(f"/api/engagements/{eid}/bank-reconciliations",
                        json={**wrong, "accounts": ["1000 Cash"]})
        assert r.status_code == 400, r.text
    r = client.post(f"/api/engagements/{eid}/bank-reconciliations",
                    json={"bank_upload_id": bank, "gl_upload_id": gl, "accounts": ["9999 Nope"]})
    assert r.json()["detail"] == "Not in the general ledger: 9999 Nope."
    rec = client.post(f"/api/engagements/{eid}/bank-reconciliations",
                      json={"bank_upload_id": bank, "gl_upload_id": gl,
                            "accounts": ["1000 Cash"]}).json()
    assert client.delete(f"/api/uploads/{bank}").status_code == 204
    assert client.get(f"/api/bank-reconciliations/{rec['id']}").status_code == 404
    log = client.get(f"/api/activity?engagement_id={eid}&action=upload.delete").json()
    assert log[0]["details"]["affects"] == ["bank"]


def test_preparers_reconcile_and_outsiders_cannot_see(real_auth):
    a = admin()
    eid, gl = setup_engagement(a)
    prep = add_user(a, "preparer")
    join(a, eid, prep)
    bank = prep.post(f"/api/engagements/{eid}/uploads", data={"kind": "bank"},
                     files={"file": (BANK.name, BANK.read_bytes(), "text/csv")}).json()["id"]
    r = prep.post(f"/api/engagements/{eid}/bank-reconciliations",
                  json={"bank_upload_id": bank, "gl_upload_id": gl, "accounts": ["1000 Cash"]})
    assert r.status_code == 201, r.text
    other = add_user(a, "reviewer")
    assert other.get(f"/api/bank-reconciliations/{r.json()['id']}").status_code == 404


# --- Roll forward -----------------------------------------------------------

def test_next_period_and_name():
    assert rollforward.next_year(date(2026, 12, 31)) == date(2027, 12, 31)
    assert rollforward.next_year(date(2027, 2, 28)) == date(2028, 2, 29)  # month end
    assert rollforward.next_year(date(2028, 2, 29)) == date(2029, 2, 28)
    assert rollforward.next_year(date(2026, 6, 15)) == date(2027, 6, 15)
    assert rollforward.next_name("FY2026") == "FY2027"
    assert rollforward.next_name("Audit 2025/2026") == "Audit 2026/2027"
    assert rollforward.next_name("Year end audit") == "Year end audit (rolled forward)"
    assert rollforward.next_name("Job 120260") == "Job 120260 (rolled forward)"


def test_roll_forward_carries_last_years_work():
    eid = engagement(with_prior=True)
    map_all(eid)
    client.put(f"/api/engagements/{eid}/materiality",
               json={"benchmark": "revenue", "benchmark_amount": 6_850_000, "benchmark_pct": 1,
                     "performance_pct": 70, "rationale": "Revenue-driven"})
    assert post_entry(eid).status_code == 201  # booked: 12,000 salaries accrued
    assert post_entry(eid, "proposed").status_code == 201  # not booked: not carried
    client.post(f"/api/engagements/{eid}/requests",
                json={"title": "Bank letters", "details": "All accounts", "due_date": "2027-01-15"})

    r = client.post(f"/api/engagements/{eid}/roll-forward", json={})
    assert r.status_code == 201, r.text
    new = r.json()
    assert new["name"] == "FY2027" and new["period_end"] == "2027-12-31"
    assert new["carried"]["mapping"] > 0 and new["carried"]["requests"] == 1
    assert new["materiality"] is None
    assert new["materiality_carried"] == {"benchmark": "revenue", "benchmark_pct": 1.0,
                                          "performance_pct": 70.0, "trivial_pct": 5.0,
                                          "rationale": "Revenue-driven"}
    [prior] = new["uploads"]
    assert prior["kind"] == "prior_tb" and prior["filename"] == "FY2026 adjusted trial balance.csv"

    # Next year's prior-year column is this year's adjusted balance.
    old = {a["account"]: a for a in client.get(f"/api/engagements/{eid}/statements")
           .json()["accounts"]}
    upload(new["id"], DATA / "sample_tb_2026.csv", "tb")
    rows = {a["account"]: a for a in client.get(f"/api/engagements/{new['id']}/statements")
            .json()["accounts"]}
    assert rows["6200"]["prior"] == old["6200"]["adjusted"] == 1_428_900 + 12_000
    assert rows["1000"]["prior"] == old["1000"]["adjusted"]
    assert all(r["line"] for r in rows.values() if r["in_tb"])  # the mapping came across
    reqs = client.get(f"/api/engagements/{new['id']}/requests").json()["requests"]
    assert [(q["number"], q["title"], q["status"], q["due_date"]) for q in reqs] == [
        (1, "Bank letters", "open", None)]

    # The old engagement is untouched, and both logs say what happened.
    assert len(client.get(f"/api/engagements/{eid}").json()["uploads"]) == 2
    old_log = client.get(f"/api/activity?engagement_id={eid}&action=engagement.roll").json()
    assert old_log[0]["summary"] == "Rolled forward to FY2027 (period end 2027-12-31)"
    new_log = [a["action"] for a in
               client.get(f"/api/activity?engagement_id={new['id']}").json()]
    assert "engagement.create" in new_log and "upload.create" in new_log

    # Rolling forward twice to the same name and period is refused.
    assert client.post(f"/api/engagements/{eid}/roll-forward", json={}).status_code == 409
    r = client.post(f"/api/engagements/{eid}/roll-forward",
                    json={"name": "FY2027 restated", "trial_balance": False, "mapping": False,
                          "materiality": False, "requests": False, "team": False})
    assert r.status_code == 201
    bare = r.json()
    assert bare["uploads"] == [] and bare["materiality_carried"] is None
    assert client.get(f"/api/engagements/{bare['id']}/requests").json()["requests"] == []


def test_roll_forward_team_and_permissions(real_auth):
    a = admin()
    eid, _ = setup_engagement(a)
    prep, rev = add_user(a, "preparer"), add_user(a, "reviewer")
    join(a, eid, prep, rev)
    assert prep.post(f"/api/engagements/{eid}/roll-forward", json={}).status_code == 403
    r = rev.post(f"/api/engagements/{eid}/roll-forward", json={})
    assert r.status_code == 201, r.text
    new = r.json()
    assert sorted(m["role"] for m in new["team"]) == ["preparer", "reviewer"]
    assert prep.get(f"/api/engagements/{new['id']}").status_code == 200
    # Without the team, whoever rolled it forward still keeps it.
    r = rev.post(f"/api/engagements/{eid}/roll-forward", json={"name": "Solo", "team": False})
    assert [m["role"] for m in r.json()["team"]] == ["reviewer"]
    assert prep.get(f"/api/engagements/{r.json()['id']}").status_code == 404
    assert "trial_balance" not in new["carried"]  # no trial balance to carry


def test_new_text_has_arabic():
    from tests.test_i18n import dictionary  # noqa: PLC0415

    d = dictionary()
    patterns = [re.compile("^" + re.sub(r"\\\{(?:t:)?\w+\\\}", ".+", re.escape(p)) + "$")
                for p in d["patterns"]]

    def translated(text):
        return text in d["exact"] or any(p.match(text) for p in patterns)

    texts = [label for _, label in bankrec.STATEMENT_ROWS]
    texts += ["Deposit in transit", "Outstanding payment", "Receipt not in the books",
              "Payment not in the books", "Bank statement", "Bank reconciliation"]
    rv, rq = {"open_notes": 0, "reviewed": 0, "total": 0}, {"overdue": 0}
    for bank in ({"reconciled": False}, {"reconciled": True, "difference": 3.0, "stale": 0},
                 {"reconciled": True, "difference": 0.0, "stale": 1},
                 {"reconciled": True, "difference": None, "stale": 2}):
        steps = dashboard.next_steps(None, {}, {"gl", "tb", "prior_tb"}, 1, {"open": 0}, 0,
                                     rv, rq, None, bank)
        texts += [s["text"] for s in steps]
    for err in ("Pick the cash account (or accounts) in the general ledger.",
                "The date window must be between 0 and 60 days.",
                "No ledger lines for those accounts up to the period end.",
                "The bank statement has no dated lines with an amount.",
                "Rolled forward to FY2027 (period end 2027-12-31)",
                "Reconciled a.csv with 1000 Cash: 3 matched, 1 outstanding, 2 not in the books",
                "This client already has an engagement called FY2027 for the period ending 2027-12-31.",
                "May be last year's item"):
        texts.append(err)
    missing = sorted({t for t in texts if not translated(t)})
    assert not missing, missing


def test_review_fixes():
    # A trial balance's Balance column is the closing balance, so it stays the amount.
    assert loader.suggest_mapping(["Account", "Debit", "Credit", "Balance"])["amount"] == "Balance"
    assert "amount" not in loader.suggest_mapping(["Date", "Debit", "Credit", "Balance"])
    # Accounts without a number are matched on their whole name.
    assert bankrec.account_code("Cash at bank") == "Cash at bank"
    assert bankrec.account_code("1000 Cash") == "1000"
    assert bankrec.cash_accounts(["Cash at bank", "Sales", "Bank - Main"], set()) == [
        "Cash at bank", "Bank - Main"]
    # Bank lines before the period are last year's; early unmatched ones get a note.
    r = reconcile(bank=[("2025-12-20", "OLD CHQ", -200.0), ("2026-01-03", "CHQ 9", -50.0),
                        ("2026-06-30", "CHARGES", -5.0), ("2026-07-01", "TRANSFER IN", 10.0)],
                  ledger=[("2026-07-01", "Receipt", 10.0)], bank_balance=0.0, book_balance=0.0)
    assert r["summary"]["bank_before_period"] == 1
    assert [(b["description"], b["early"]) for b in r["bank_only"]] == [
        ("CHQ 9", True), ("CHARGES", False)]
    assert r["settings"]["period_start"] == "2026-01-01"


def test_bank_section_is_only_needed_once_there_is_bank_work():
    eid = engagement(with_prior=False)
    rv = client.get(f"/api/engagements/{eid}/review").json()
    bank = next(s for s in rv["sections"] if s["id"] == "bank")
    assert not bank["needed"] and rv["total"] == 7
    upload(eid, BANK, "bank")
    rv = client.get(f"/api/engagements/{eid}/review").json()
    assert next(s for s in rv["sections"] if s["id"] == "bank")["needed"] and rv["total"] == 8


def test_roll_forward_after_a_short_first_period():
    c = client.post("/api/clients", json={"name": "Startup"}).json()
    e = client.post(f"/api/clients/{c['id']}/engagements",
                    json={"name": "FY2025", "period_start": "2025-06-15",
                          "period_end": "2025-12-31"}).json()
    new = client.post(f"/api/engagements/{e['id']}/roll-forward", json={}).json()
    assert (new["period_start"], new["period_end"]) == ("2026-01-01", "2026-12-31")
