"""Lead sheets, adjusting entries and draft financial statements."""

import io
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app import statements
from app.main import app

client = TestClient(app)
DATA = Path(__file__).resolve().parent.parent / "sample_data"


def tb_frame(rows):
    return pd.DataFrame([{"account": a, "description": n, "amount": amt} for a, n, amt in rows])


TB = tb_frame([("1000", "Cash", 500.0), ("1200", "Receivables", 300.0),
               ("2000", "Payables", -200.0), ("3000", "Share capital", -100.0),
               ("4000", "Sales", -1000.0), ("6100", "Salaries", 500.0)])
MAPPING = {"1000": "cash", "1200": "receivables", "2000": "payables", "3000": "share_capital",
           "4000": "revenue", "6100": "staff_costs"}


def entry(kind, lines, ref="AJE-1", id_=1):
    return {"id": id_, "kind": kind, "ref": ref, "description": "x",
            "lines": [{"account": a, "debit": d, "credit": c} for a, d, c in lines]}


def test_suggest_line_uses_names_then_account_numbers():
    s = statements.suggest_line
    assert s("2500", "Bank Loan") == "borrowings"  # not cash, despite "bank"
    assert s("1010", "Bank - current account") == "cash"
    assert s("1600", "Accumulated depreciation - vehicles") == "ppe"
    assert s("6800", "Depreciation") == "depreciation"
    assert s("4000", "Sales") == "revenue"
    assert s("4900", "Interest income") == "other_income"
    assert s("5999", "Misc") == "cost_of_sales"  # by account number
    assert s("ABC", "Misc") is None


def test_entry_lines_must_balance():
    ok = [{"account": "1", "debit": 10, "credit": 0}, {"account": "2", "debit": 0, "credit": 10}]
    statements.check_entry_lines(ok)
    with pytest.raises(statements.StatementsError, match="must be equal"):
        statements.check_entry_lines([ok[0], {"account": "2", "credit": 9}])
    with pytest.raises(statements.StatementsError, match="either a debit or a credit"):
        statements.check_entry_lines([ok[0], {"account": "2", "debit": 5, "credit": 5}])
    with pytest.raises(statements.StatementsError, match="at least two"):
        statements.check_entry_lines([ok[0]])


def test_statements_balance_and_booked_entries_adjust():
    aje = entry("adjusting", [("6100", 50, 0), ("2000", 0, 50)])  # unrecorded salary accrual
    paje = entry("proposed", [("1200", 0, 40), ("4000", 40, 0)], "PAJE-1", 2)  # cut-off
    out = statements.build(TB, None, MAPPING, [aje, paje])
    st = out["statements"]
    assert st["difference"]["current"] == 0
    assert st["profit"]["current"] == 1000 - 500 - 50  # AJE booked, PAJE not
    payables = next(s for s in out["lead_sheets"] if s["line"] == "payables")
    assert payables["totals"] == {"unadjusted": 200, "adjusting": 50, "reclassification": 0,
                                  "adjusted": 250, "prior": 0}
    unc = out["uncorrected"]
    assert unc["profit"] == -40 and unc["net_assets"] == -40
    assert unc["conclusion"] is None  # no materiality set


def test_unmapped_accounts_are_listed_and_unbalance_the_statements():
    mapping = {k: v for k, v in MAPPING.items() if k != "1200"}
    out = statements.build(TB, None, mapping, [])
    assert [u["account"] for u in out["unmapped"]] == ["1200"]
    assert out["unmapped"][0]["suggested"] == "receivables"
    assert out["statements"]["difference"]["current"] == -300


def test_uncorrected_misstatements_against_materiality():
    paje = entry("proposed", [("4000", 900, 0), ("1200", 0, 900)], "PAJE-1")
    mat = {"overall_materiality": 1000, "performance_materiality": 750,
           "trivial_threshold": 50}
    unc = statements.build(TB, None, MAPPING, [paje], mat)["uncorrected"]
    assert unc["pct_of_overall"] == 90.0
    assert "above performance materiality" in unc["conclusion"]
    paje["lines"][0]["debit"] = paje["lines"][1]["credit"] = 1200
    unc = statements.build(TB, None, MAPPING, [paje], mat)["uncorrected"]
    assert "material" in unc["conclusion"] and unc["pct_of_overall"] == 120.0


# --- API ----------------------------------------------------------------------


def engagement(with_prior=True, client_name="Acme"):
    c = client.post("/api/clients", json={"name": client_name}).json()
    e = client.post(f"/api/clients/{c['id']}/engagements",
                    json={"name": "FY2026", "period_end": "2026-12-31"}).json()
    files = [("sample_tb_2026.csv", "tb")] + ([("sample_tb_2025.csv", "prior_tb")]
                                               if with_prior else [])
    for name, kind in files:
        r = client.post(f"/api/engagements/{e['id']}/uploads", data={"kind": kind},
                        files={"file": (name, (DATA / name).read_bytes(), "text/csv")})
        assert r.status_code == 201, r.text
    return e["id"]


def map_all(eid):
    data = client.get(f"/api/engagements/{eid}/statements").json()
    r = client.put(f"/api/engagements/{eid}/mapping",
                   json={"mappings": {a["account"]: a["suggested"] for a in data["accounts"]}})
    assert r.status_code == 200
    return r.json()


def test_no_trial_balance_yet():
    c = client.post("/api/clients", json={"name": "Empty"}).json()
    e = client.post(f"/api/clients/{c['id']}/engagements",
                    json={"name": "FY", "period_end": "2026-12-31"}).json()
    data = client.get(f"/api/engagements/{e['id']}/statements").json()
    assert data["tb"] is None and data["entries"] == []
    assert client.get(f"/api/engagements/{e['id']}/statements/export").status_code == 400


def test_mapping_and_statements_from_sample_files():
    eid = engagement()
    assert map_all(eid)["changed"] == 17
    data = client.get(f"/api/engagements/{eid}/statements").json()
    assert data["unmapped"] == [] and data["has_prior"]
    assert data["statements"]["difference"] == {"current": 0, "prior": 0}
    assert data["statements"]["profit"] == {"current": 1060840, "prior": 624500}
    bad = client.put(f"/api/engagements/{eid}/mapping", json={"mappings": {"1000": "nope"}})
    assert bad.status_code == 400
    client.put(f"/api/engagements/{eid}/mapping", json={"mappings": {"1000": None}})
    data = client.get(f"/api/engagements/{eid}/statements").json()
    assert [u["account"] for u in data["unmapped"]] == ["1000"]
    log = client.get("/api/activity", params={"engagement_id": eid}).json()
    assert log[0]["action"] == "mapping.update" and log[0]["details"]["mappings"] == {"1000": None}


def post_entry(eid, kind="adjusting", lines=None, **extra):
    lines = lines or [{"account": "6200", "debit": 12000}, {"account": "2100", "credit": 12000}]
    return client.post(f"/api/engagements/{eid}/adjusting-entries",
                       json={"kind": kind, "description": "Unrecorded payroll accrual",
                             "lines": lines, **extra})


def test_adjusting_entries_numbering_and_effect():
    eid = engagement()
    map_all(eid)
    a1 = post_entry(eid).json()
    a2 = post_entry(eid).json()
    p1 = post_entry(eid, "proposed").json()
    assert [a1["ref"], a2["ref"], p1["ref"]] == ["AJE-1", "AJE-2", "PAJE-1"]
    data = client.get(f"/api/engagements/{eid}/statements").json()
    assert data["statements"]["profit"]["current"] == 1060840 - 24000
    assert data["statements"]["difference"]["current"] == 0
    assert data["uncorrected"]["profit"] == -12000
    booked = client.patch(f"/api/adjusting-entries/{p1['id']}", json={"kind": "adjusting"})
    assert booked.json()["ref"] == "AJE-3"
    data = client.get(f"/api/engagements/{eid}/statements").json()
    assert data["statements"]["profit"]["current"] == 1060840 - 36000
    assert data["uncorrected"]["entries"] == []
    assert client.delete(f"/api/adjusting-entries/{a1['id']}").status_code == 204
    assert post_entry(eid).json()["ref"] == "AJE-4"  # numbers are not reused
    actions = [x["action"] for x in client.get("/api/activity",
                                               params={"engagement_id": eid}).json()]
    assert actions[:3] == ["entry.create", "entry.delete", "entry.update"]


def test_entries_are_validated():
    eid = engagement(with_prior=False)
    unbalanced = [{"account": "6200", "debit": 100}, {"account": "2100", "credit": 90}]
    r = post_entry(eid, lines=unbalanced)
    assert r.status_code == 400 and "must be equal" in r.json()["detail"]
    assert post_entry(eid, lines=[{"account": "6200", "debit": 1}]).status_code == 422
    assert post_entry(eid, lines=[{"account": " ", "debit": 1},
                                  {"account": "2", "credit": 1}]).status_code == 422
    assert post_entry(eid, lines=[{"account": "1", "debit": -1},
                                  {"account": "2", "credit": -1}]).status_code == 422
    assert post_entry(eid, kind="other").status_code == 422
    r = client.post(f"/api/engagements/{eid}/adjusting-entries",
                    json={"description": "  ", "lines": unbalanced})
    assert r.status_code == 422


def test_export_is_safe_and_logged():
    eid = engagement(client_name="北京贸易")
    map_all(eid)
    post_entry(eid, "proposed", lines=[{"account": "4000", "debit": 5, "memo": "=HYPERLINK(1)"},
                                       {"account": "1200", "credit": 5}],
               date="2026-12-31")
    r = client.get(f"/api/engagements/{eid}/statements/export")
    assert r.status_code == 200 and "filename*=UTF-8''" in r.headers["content-disposition"]
    book = openpyxl.load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Summary", "Statements", "Ratios", "Lead sheets", "Adjusted TB",
                               "Adjusting entries", "Uncorrected", "Cash flow"]
    cells = [c for ws in book.worksheets for row in ws.iter_rows() for c in row]
    assert not [c.value for c in cells if c.data_type == "f"]
    assert "=HYPERLINK(1)" in [c.value for c in cells]
    log = client.get("/api/activity", params={"engagement_id": eid}).json()[0]
    assert log["action"] == "statements.download" and len(log["details"]["sha256"]) == 64
    assert log["details"]["entries"] == ["PAJE-1"]


# --- Review fixes ---------------------------------------------------------------


def test_suggestions_match_whole_words_in_the_right_order():
    s = statements.suggest_line
    assert s("1300", "Other current assets") == "other_current_assets"  # not "rent"
    assert s("2600", "Current portion of long-term debt") == "borrowings"
    assert s("3000", "Common stock") == "share_capital"  # not inventory
    assert s("1250", "Loan receivable") == "receivables"
    assert s("2150", "Customer deposits") == "accruals"
    assert s("1260", "Accrued income") == "other_current_assets"
    assert s("2160", "Payroll liabilities") == "accruals"
    assert s("Cash at bank", None) == "cash"  # account names with no description column


def test_unmapped_proposed_entry_accounts_block_the_conclusion():
    paje = entry("proposed", [("X1", 500, 0), ("X2", 0, 500)], "PAJE-1")
    mat = {"overall_materiality": 100, "performance_materiality": 75, "trivial_threshold": 5}
    unc = statements.build(TB, None, MAPPING, [paje], mat)["uncorrected"]
    assert unc["unclassified"] == ["X1", "X2"]
    assert unc["conclusion"].startswith("Map X1, X2") and "trivial" not in unc["conclusion"]


def test_entry_totals_must_be_finite_and_balance_to_the_cent():
    big = [{"account": "1", "debit": 1.7e308}, {"account": "1", "debit": 1.7e308},
           {"account": "2", "credit": 1.7e308}, {"account": "2", "credit": 1.7e308},
           {"account": "2", "credit": 1.7e308}]
    with pytest.raises(statements.StatementsError, match="too large"):
        statements.check_entry_lines(big)
    eid = engagement(with_prior=False)
    assert post_entry(eid, lines=[{"account": "1", "debit": 1e14},
                                  {"account": "2", "credit": 1e14}]).status_code == 422
    r = post_entry(eid, lines=[{"account": "6200", "debit": 100.004},
                               {"account": "2100", "credit": 100}])
    assert r.status_code == 201 and r.json()["lines"][0]["debit"] == 100.0


def test_blank_and_numeric_account_codes_are_normalized():
    from app import loader
    csv = b"account,description,amount\n1000,Cash,5\n,Misc,1\n   ,Petty,2\n2000,Payables,-8\n"
    df = loader.load("tb.csv", csv).df
    balances = statements._balances(df)
    assert list(balances) == ["1000", "(blank)", "2000"]
    assert balances["(blank)"] == {"balance": 3.0, "name": "Misc"}


def test_export_leaves_prior_year_blank_without_a_prior_tb():
    eid = engagement(with_prior=False)
    map_all(eid)
    book = openpyxl.load_workbook(io.BytesIO(
        client.get(f"/api/engagements/{eid}/statements/export").content))
    ws = book["Statements"]
    assert ws["D1"].value == "Prior" and all(c.value is None for c in ws["D"][1:])
    assert any(isinstance(c.value, (int, float)) for c in ws["C"][1:])


def test_entry_numbers_are_never_reused():
    eid = engagement(with_prior=False)
    a1, a2 = post_entry(eid).json(), post_entry(eid).json()
    client.delete(f"/api/adjusting-entries/{a2['id']}")
    assert post_entry(eid).json()["ref"] == "AJE-3"
    client.patch(f"/api/adjusting-entries/{a1['id']}", json={"kind": "proposed"})
    assert post_entry(eid).json()["ref"] == "AJE-4"


def _race(monkeypatch, call, n=2):
    import threading
    import time
    from concurrent.futures import ThreadPoolExecutor

    from app import leadsheets
    real = leadsheets.log

    def slow_log(*args, **kwargs):  # widen the window between reading and committing
        time.sleep(0.2)
        return real(*args, **kwargs)

    monkeypatch.setattr(leadsheets, "log", slow_log)
    barrier = threading.Barrier(n)
    quiet = TestClient(app, raise_server_exceptions=False)

    def run(_):
        barrier.wait()
        return call(quiet)

    with ThreadPoolExecutor(n) as pool:
        return list(pool.map(run, range(n)))


def test_concurrent_mapping_saves_do_not_fail(monkeypatch):
    eid = engagement(with_prior=False)
    codes = _race(monkeypatch, lambda c: c.put(f"/api/engagements/{eid}/mapping",
                                               json={"mappings": {"1000": "cash"}}).status_code)
    assert sorted(codes) == [200, 200]


def test_concurrent_mark_booked_books_once(monkeypatch):
    eid = engagement(with_prior=False)
    p1 = post_entry(eid, "proposed").json()
    refs = _race(monkeypatch, lambda c: c.patch(f"/api/adjusting-entries/{p1['id']}",
                                                json={"kind": "adjusting"}).json()["ref"])
    assert refs == ["AJE-1", "AJE-1"]
    log = client.get("/api/activity", params={"engagement_id": eid}).json()
    assert [x["action"] for x in log].count("entry.update") == 1


# --- Copilot review fixes -------------------------------------------------------


def test_keywords_need_the_whole_word():
    s = statements.suggest_line
    assert s("3100", "Stockholders' equity") == "retained_earnings"  # not inventory
    assert s("6900", "Bankruptcy costs") == "other_expenses"  # not cash
    assert s("1400", "Inventories") == "inventory"  # stems still match
    assert s("6400", "Administrative expenses") == "admin"


def test_only_saved_mappings_classify_misstatements():
    paje = entry("proposed", [("6100", 5, 0), ("1200", 0, 5)], "PAJE-1")
    mat = {"overall_materiality": 1000, "performance_materiality": 750, "trivial_threshold": 50}
    mapping = {k: v for k, v in MAPPING.items() if k != "6100"}  # suggested, but not saved
    unc = statements.build(TB, None, mapping, [paje], mat)["uncorrected"]
    assert unc["unclassified"] == ["6100"] and unc["conclusion"].startswith("Map 6100")
    unc = statements.build(TB, None, MAPPING, [paje], mat)["uncorrected"]
    assert unc["unclassified"] == [] and "trivial" in unc["conclusion"]


def test_long_entries_of_large_amounts_balance_exactly():
    lines = [{"account": "1", "debit": 9999999999999.99}] * 100 + \
            [{"account": "2", "credit": 9999999999999.99}] * 100
    statements.check_entry_lines(lines)
    with pytest.raises(statements.StatementsError, match="must be equal"):
        statements.check_entry_lines(lines[:-1] + [{"account": "2", "credit": 9999999999999.98}])


def test_numbers_from_before_the_counter_are_not_reused():
    from app import db
    eid = engagement(with_prior=False)
    post_entry(eid)
    a2 = post_entry(eid).json()
    with db.SessionLocal() as s:  # a database from before entry numbers had counters
        s.query(db.Guard).filter(db.Guard.key.like("entry-number:%")).delete(
            synchronize_session=False)
        s.commit()
    client.delete(f"/api/adjusting-entries/{a2['id']}")
    assert post_entry(eid).json()["ref"] == "AJE-3"
