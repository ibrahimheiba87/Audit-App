"""Ratio and trend analysis on the draft statements, and the engagement overview."""

import io

import openpyxl
import pytest

from app import dashboard, ratios, statements
from tests.test_statements import DATA, TB, client, engagement, map_all, post_entry, tb_frame

# A small company: 1,000 revenue, 600 cost of sales, 200 overheads, 50 interest, 30 tax.
ROWS = [("1000", "Cash", 150.0), ("1200", "Receivables", 250.0), ("1300", "Inventory", 120.0),
        ("1500", "Equipment", 480.0), ("2000", "Payables", -90.0), ("2100", "Accruals", -110.0),
        ("2500", "Bank loan", -300.0), ("3000", "Share capital", -100.0),
        ("3100", "Retained earnings", -280.0), ("4000", "Sales", -1000.0),
        ("5000", "Purchases", 600.0), ("6100", "Salaries", 200.0), ("7000", "Interest", 50.0),
        ("8000", "Income tax", 30.0)]
MAPPING = {"1000": "cash", "1200": "receivables", "1300": "inventory", "1500": "ppe",
           "2000": "payables", "2100": "accruals", "2500": "borrowings",
           "3000": "share_capital", "3100": "retained_earnings", "4000": "revenue",
           "5000": "cost_of_sales", "6100": "staff_costs", "7000": "finance_costs",
           "8000": "income_tax"}


def ratio_table(prior_rows=None):
    prior = tb_frame(prior_rows) if prior_rows else None
    data = statements.build(tb_frame(ROWS), prior, MAPPING, [])
    out = ratios.compute(data)
    return out, {r["id"]: r for r in out["ratios"]}


def test_ratios_from_the_statements():
    out, r = ratio_table()
    assert not out["has_prior"] and out["flagged"] == 0
    assert r["current"]["current"] == round(520 / 200, 2)
    assert r["quick"]["current"] == 2.0
    assert r["gross_margin"]["current"] == 40.0  # cost of sales counted once, not twice
    assert r["operating_margin"]["current"] == 20.0
    assert r["net_margin"]["current"] == 12.0
    assert r["receivable_days"]["current"] == round(250 / 1000 * 365, 2)
    assert r["inventory_days"]["current"] == round(120 / 600 * 365, 2)
    assert r["debt_to_equity"]["current"] == round(300 / (380 + 120), 2)
    assert r["interest_cover"]["current"] == 4.0
    assert r["gross_margin"]["prior"] is None and r["gross_margin"]["flag"] is None
    assert out["warnings"] == []


def test_big_moves_on_last_year_are_flagged():
    prior = [(a, n, amt) for a, n, amt in ROWS]
    prior[1] = ("1200", "Receivables", 150.0)  # receivable days up from 55 to 91 days
    prior[2] = ("1300", "Inventory", 110.0)    # inventory days up a little
    prior[0] = ("1000", "Cash", 260.0)
    out, r = ratio_table(prior)
    assert out["has_prior"]
    assert r["receivable_days"]["flag"] == "Up 67% on last year"
    assert r["inventory_days"]["flag"] is None
    assert r["gross_margin"]["change"] == 0 and r["gross_margin"]["flag"] is None
    prior = [(a, n, -900.0 if a == "4000" else amt) for a, n, amt in ROWS]  # lower sales
    prior = [(a, n, amt) if a != "3100" else (a, n, -180.0) for a, n, amt in prior]
    _, r = ratio_table(prior)
    assert r["gross_margin"]["prior"] == round(300 / 900 * 100, 2)
    assert r["gross_margin"]["flag"] == "Up 6.7 points on last year"
    assert ratio_table(prior)[0]["flagged"] == sum(1 for x in r.values() if x["flag"])


def test_going_concern_warnings():
    data = statements.build(TB, None, {"1000": "cash", "1200": "receivables",
                                       "2000": "borrowings_current", "3000": "share_capital",
                                       "4000": "revenue", "6100": "staff_costs"}, [])
    assert ratios.compute(data)["warnings"] == []
    loss = [("1000", "Cash", 100.0), ("2000", "Payables", -300.0), ("3000", "Capital", -100.0),
            ("4000", "Sales", -500.0), ("5000", "Purchases", 600.0), ("7000", "Interest", 200.0)]
    data = statements.build(tb_frame(loss), None, {
        "1000": "cash", "2000": "payables", "3000": "share_capital", "4000": "revenue",
        "5000": "cost_of_sales", "7000": "finance_costs"}, [])
    out = ratios.compute(data)
    assert [w["id"] for w in out["warnings"]] == [
        "net_current_liabilities", "negative_equity", "loss", "low_interest_cover",
        "negative_gross_margin"]
    by_id = {r["id"]: r for r in out["ratios"]}
    assert by_id["debt_to_equity"]["current"] is None  # meaningless with negative equity


@pytest.mark.parametrize("value", [0.0, 0.001])
def test_no_division_by_zero(value):
    rows = [("1000", "Cash", value), ("3000", "Capital", -value)]
    data = statements.build(tb_frame(rows), None, {"1000": "cash", "3000": "share_capital"}, [])
    out = ratios.compute(data)
    assert all(r["current"] is None for r in out["ratios"])


def test_statements_and_workpapers_carry_the_ratios():
    eid = engagement()
    map_all(eid)
    data = client.get(f"/api/engagements/{eid}/statements").json()
    r = {x["id"]: x for x in data["ratios"]["ratios"]}
    assert data["ratios"]["has_prior"] and r["current"]["prior"] is not None
    assert data["ratios"]["flagged"] == sum(1 for x in r.values() if x["flag"]) > 0
    book = openpyxl.load_workbook(io.BytesIO(
        client.get(f"/api/engagements/{eid}/statements/export").content))
    assert book.sheetnames[:3] == ["Summary", "Statements", "Ratios"]
    rows = list(book["Ratios"].iter_rows(values_only=True))
    assert rows[0][:3] == ("Group", "Ratio", "Unit")
    assert {row[1] for row in rows[1:]} >= {"Current ratio", "Gross margin", "Receivable days"}


def test_dashboard_lists_the_next_steps():
    c = client.post("/api/clients", json={"name": "New Co"}).json()
    e = client.post(f"/api/clients/{c['id']}/engagements",
                    json={"name": "FY", "period_end": "2025-12-31"}).json()
    d = client.get(f"/api/engagements/{e['id']}/dashboard").json()
    assert d["statements"] is None and d["days_since_period_end"] > 0
    texts = [s["text"] for s in d["next_steps"]]
    assert texts[:2] == ["Set materiality.", "Upload this year's trial balance."]
    assert [s["priority"] for s in d["next_steps"]] == sorted(
        (s["priority"] for s in d["next_steps"]), key=["high", "medium", "low"].index)


def test_dashboard_follows_the_work():
    eid = engagement()
    d = client.get(f"/api/engagements/{eid}/dashboard").json()
    assert d["upload_kinds"] == ["prior_tb", "tb"]
    assert d["statements"]["unmapped"] == 17
    assert "Map 17 accounts to a statement line." in [s["text"] for s in d["next_steps"]]
    map_all(eid)
    r = client.put(f"/api/engagements/{eid}/materiality",
                   json={"benchmark": "revenue", "benchmark_amount": 100000, "benchmark_pct": 1})
    assert r.status_code == 200, r.text
    post_entry(eid, "proposed", lines=[{"account": "4000", "debit": 5000},
                                       {"account": "1200", "credit": 5000}])
    d = client.get(f"/api/engagements/{eid}/dashboard").json()
    steps = {s["text"]: s for s in d["next_steps"]}
    assert d["materiality"]["overall_materiality"] == 1000
    assert "Set materiality." not in steps and d["statements"]["unmapped"] == 0
    assert d["statements"]["uncorrected_entries"] == 1
    material = d["statements"]["uncorrected"]
    assert material.startswith("Uncorrected misstatements are material")
    assert steps[material]["priority"] == "high" and steps[material]["panel"] == "fs-panel"
    assert any(t.startswith("Explain ") for t in steps)



def test_flags_compare_against_the_size_of_last_year():
    prior = [(a, n, -2000.0 if a == "4000" else amt) for a, n, amt in ROWS]
    prior = [(a, n, amt) if a != "3100" else (a, n, 720.0) for a, n, amt in prior]
    _, r = ratio_table(prior)  # last year: much higher sales, so margins and cover moved
    assert r["interest_cover"]["prior"] == round(1200 / 50, 2)
    assert r["interest_cover"]["flag"] == "Down 83% on last year"
    assert r["gross_margin"]["flag"].startswith("Down ")


def test_only_general_ledger_runs_count_as_running_the_checks():
    eid = engagement(with_prior=False)
    gl = client.post(f"/api/engagements/{eid}/uploads", data={"kind": "gl"},
                     files={"file": ("gl.csv", (DATA / "sample_gl.csv").read_bytes(), "text/csv")})
    assert gl.status_code == 201, gl.text
    step = "Run the audit checks on the general ledger."
    texts = lambda: [s["text"] for s in  # noqa: E731
                     client.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]]
    assert step in texts()
    tb = next(u for u in client.get(f"/api/engagements/{eid}").json()["uploads"]
              if u["kind"] == "tb")
    assert client.post(f"/api/engagements/{eid}/runs", json={"upload_id": tb["id"]}).status_code == 201
    assert step in texts()
    r = client.post(f"/api/engagements/{eid}/runs", json={"upload_id": gl.json()["id"]})
    assert r.status_code == 201 and step not in texts()


def test_dashboard_uses_the_chosen_trial_balances():
    eid = engagement()
    map_all(eid)
    prior = next(u for u in client.get(f"/api/engagements/{eid}").json()["uploads"]
                 if u["kind"] == "prior_tb")
    latest = client.get(f"/api/engagements/{eid}/dashboard").json()["statements"]
    chosen = client.get(f"/api/engagements/{eid}/dashboard",
                        params={"prior_upload_id": prior["id"]}).json()["statements"]
    assert chosen == latest  # only one prior-year file, so the same answer
    steps = client.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]
    assert {s["tab"] for s in steps if s["panel"] == "fs-panel"} <= {"ratios", "mapping",
                                                                     "uncorrected", None}


def test_every_next_step_and_flag_has_an_arabic_translation():
    import re  # noqa: PLC0415

    from tests.test_i18n import dictionary  # noqa: PLC0415

    d = dictionary()
    patterns = [re.compile("^" + re.sub(r"\\\{(?:t:)?\w+\\\}", ".+", re.escape(p)) + "$")
                for p in d["patterns"]]
    translated = lambda text: text in d["exact"] or any(p.match(text) for p in patterns)  # noqa: E731
    texts = []
    for rows in (ROWS, [("1000", "Cash", 100.0), ("2000", "Payables", -300.0),
                        ("3000", "Capital", -100.0), ("4000", "Sales", -500.0),
                        ("5000", "Purchases", 600.0), ("7000", "Interest", 200.0)]):
        mapping = {**MAPPING, "3000": "share_capital"}
        out = ratios.compute(statements.build(tb_frame(rows), tb_frame(ROWS), mapping, []))
        texts += [w["text"] for w in out["warnings"]] + [r["flag"] for r in out["ratios"] if r["flag"]]
        texts += [r["label"] for r in out["ratios"]] + [r["basis"] for r in out["ratios"]]
        texts += [r["group_label"] for r in out["ratios"]]
    for n in (1, 2):
        rv = {"open_notes": n, "reviewed": 0, "total": n}
        rq = {"overdue": n}
        fs = {"unmapped": n, "difference": 1, "warnings": [], "ratio_flags": n,
              "uncorrected": None}
        steps = dashboard.next_steps(None, None, {"gl"}, 0, {"open": n}, n, rv, rq, fs)
        steps += dashboard.next_steps(None, None, set(), 0, {"open": n}, 0, rv, rq,
                                      {"error": "x"})
        texts += [s["text"] for s in steps]
    missing = sorted({t for t in texts if not translated(t)})
    assert not missing, missing
