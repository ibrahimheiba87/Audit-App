"""AI risk assessment and management letter, with the Claude API faked: no real calls."""

import io
import json

import openpyxl
import pytest

from app import ai_writer, assistant
from tests.test_ai_tools import activity, fake_api, with_gl  # noqa: F401 (fixture)
from tests.test_statements import client, engagement, map_all

RISKS = {"summary": "Revenue grew fast.", "risks": [
    {"area": "receivables", "title": "Receivables may be overstated", "level": "higher",
     "assertions": ["existence", "valuation", "made_up"], "why": "Up 40% on last year.",
     "procedures": ["Circularise the largest balances", 7, None, {}]},
    {"area": "revenue", "title": "=HYPERLINK(\"x\")", "level": "significant",
     "assertions": ["occurrence", "cut_off", "occurrence"], "why": "Presumed fraud risk.",
     "procedures": ["Test cut-off"]},
    {"area": "nowhere", "title": "Override", "level": "extreme", "assertions": "x"},
    {"area": "cash", "title": ""}, {"title": "Odd level", "level": ["high"], "area": {}},
    "junk"]}


def findings_engagement():
    eid = engagement()
    gl = with_gl(eid)
    assert client.post(f"/api/engagements/{eid}/runs", json={"upload_id": gl}).status_code == 201
    return eid, client.get(f"/api/engagements/{eid}/findings").json()["findings"]


def test_off_without_a_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    eid = engagement()
    for path in ("risk-assessment", "management-letter"):
        r = client.post(f"/api/engagements/{eid}/ai/{path}", json={})
        assert r.status_code == 503 and "ANTHROPIC_API_KEY" in r.json()["detail"]
    assert client.get(f"/api/engagements/{eid}/ai-drafts").json() == {
        "risk": {"latest": None, "count": 0}, "letter": {"latest": None, "count": 0}}


def test_risk_assessment_is_checked_saved_and_logged(fake_api):  # noqa: F811
    eid = engagement()
    map_all(eid)
    fake_api["answers"].append(RISKS)
    r = client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={"language": "ar"})
    assert r.status_code == 201, r.text
    d = r.json()
    risks = d["content"]["risks"]
    # Significant first; unknown areas become the whole statements; empty titles are dropped.
    assert [(x["area"], x["level"]) for x in risks] == [
        ("revenue", "significant"), ("receivables", "higher"), ("entity", "higher"),
        ("entity", "higher")]
    assert risks[0]["assertions"] == ["occurrence", "cut_off"]
    assert risks[1]["assertions"] == ["existence", "valuation"]
    assert risks[1]["procedures"] == ["Circularise the largest balances", "7"]
    assert risks[2]["area_label"] == "Whole financial statements"
    system, prompt, kw = fake_api["calls"][0]
    assert system == ai_writer.RISK_SYSTEM and "Modern Standard Arabic" in prompt
    context = json.loads(prompt.split("<engagement>\n")[1].split("\n</engagement>")[0])
    assert {x["line"] for x in context["statement_lines"]} >= {"revenue", "cash"}
    assert context["engagement"]["framework"] == "IFRS" and context["ratios"]
    assert kw["max_tokens"] == ai_writer.MAX_TOKENS
    assert d["sent"]["has_prior"] and d["language"] == "ar"
    [entry] = activity(eid, "assistant.risk_assessment")
    assert entry["summary"] == "Asked the AI assistant for a risk assessment"
    assert entry["details"]["outcome"] == "ok"
    drafts = client.get(f"/api/engagements/{eid}/ai-drafts").json()
    assert drafts["risk"]["latest"]["id"] == d["id"] and drafts["risk"]["count"] == 1
    # The Excel download: formula-looking text stays text, and it's logged.
    r = client.get(f"/api/ai-drafts/{d['id']}/export")
    book = openpyxl.load_workbook(io.BytesIO(r.content))
    cells = [c for row in book["Risks"].iter_rows() for c in row]
    assert any(c.value == '=HYPERLINK("x")' and c.data_type == "s" for c in cells)
    assert activity(eid, "ai_draft.download")[0]["summary"] == "Downloaded AI risk assessment #1"


def test_risk_assessment_needs_a_trial_balance_and_some_risks(fake_api):  # noqa: F811
    c = client.post("/api/clients", json={"name": "Empty"}).json()
    e = client.post(f"/api/clients/{c['id']}/engagements",
                    json={"name": "FY", "period_end": "2026-12-31"}).json()
    r = client.post(f"/api/engagements/{e['id']}/ai/risk-assessment", json={})
    assert r.status_code == 400 and not fake_api["calls"]
    eid = engagement()
    fake_api["answers"] += [{"risks": []}, "not json"]
    for _ in range(2):
        assert client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={}).status_code == 502
    assert client.get(f"/api/engagements/{eid}/ai-drafts").json()["risk"]["count"] == 0
    # An answer that can't be used is logged as a failure, not as a success.
    entries = activity(eid, "assistant.risk_assessment")
    assert {e["summary"] for e in entries} == {"Asked the AI assistant for a risk assessment (failed)"}
    assert "no risks" in entries[-1]["details"]["outcome"]


def test_management_letter_cites_only_findings_it_was_sent(fake_api):  # noqa: F811
    eid, findings = findings_engagement()
    ids = [f["id"] for f in findings]
    client.patch(f"/api/findings/{ids[0]}", json={"status": "cleared", "note": "Explained"})
    fake_api["answers"].append({"intro": "Dear management,", "closing": "Yours,", "points": [
        {"title": "Journal approval", "finding_ids": [ids[1], ids[1], ids[0], 999999, True, "x"],
         "observation": "Entries posted at weekends.", "implication": "Errors may go unseen.",
         "recommendation": "Review weekend entries.", "priority": "urgent"},
        {"title": "", "finding_ids": [ids[2]]},
        {"title": "Odd shapes", "finding_ids": [[1], {}], "priority": {}}, "junk"]})
    r = client.post(f"/api/engagements/{eid}/ai/management-letter", json={})
    assert r.status_code == 201, r.text
    d = r.json()
    point, odd = d["content"]["points"]
    assert point["finding_ids"] == [ids[1]] and point["priority"] == "medium"
    assert odd["finding_ids"] == [] and odd["priority"] == "medium"
    prompt = fake_api["calls"][0][1]
    sent = [json.loads(x) for x in prompt.split("<findings>\n")[1].split("\n</findings>")[0]
            .splitlines()]
    assert ids[0] not in {f["number"] for f in sent}  # cleared isn't included by default
    assert d["sent"]["statuses"] == ["open", "reviewed"]
    [entry] = activity(eid, "assistant.management_letter")
    assert entry["summary"] == (f"Asked the AI assistant to draft a management letter from "
                                f"{len(sent)} findings")
    book = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/ai-drafts/{d['id']}/export").content))
    assert [c.value for c in book["Letter points"][2]][-1] == f"#{ids[1]}"


def test_management_letter_needs_findings(fake_api):  # noqa: F811
    eid = engagement()
    r = client.post(f"/api/engagements/{eid}/ai/management-letter",
                    json={"statuses": ["cleared"]})
    assert r.status_code == 400 and "no findings" in r.json()["detail"]
    assert client.post(f"/api/engagements/{eid}/ai/management-letter",
                       json={"statuses": []}).status_code == 422
    assert not fake_api["calls"]


def test_long_registers_send_the_most_severe(fake_api, monkeypatch):  # noqa: F811
    monkeypatch.setattr(ai_writer, "MAX_LETTER_FINDINGS", 3)
    eid, findings = findings_engagement()
    fake_api["answers"].append({"points": [{"title": "x", "finding_ids": []}]})
    d = client.post(f"/api/engagements/{eid}/ai/management-letter", json={}).json()
    assert d["sent"] == {"findings": 3, "findings_total": len(findings),
                         "statuses": ["open", "reviewed"]}
    worst = min(ai_writer.SEVERITY_RANK[f["severity"]] for f in findings)
    prompt = fake_api["calls"][0][1]
    first = json.loads(prompt.split("<findings>\n")[1].splitlines()[0])
    assert ai_writer.SEVERITY_RANK[first["severity"]] == worst
    # The auditor can see that the letter doesn't cover every finding.
    book = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/ai-drafts/{d['id']}/export").content))
    summary = {r[0].value: r[1].value for r in book["Summary"].iter_rows()}
    assert summary["Coverage"].startswith(f"Written from 3 of {len(findings)} findings")


def test_risk_assessment_sees_findings_added_by_hand(fake_api):  # noqa: F811
    eid, _ = findings_engagement()
    r = client.post(f"/api/engagements/{eid}/findings", json={
        "severity": "low", "detail": "Inventory count not attended", "account": "1300",
        "amount": 5000})
    assert r.status_code == 201, r.text
    fake_api["answers"].append(RISKS)
    assert client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={}).status_code == 201
    prompt = fake_api["calls"][0][1]
    context = json.loads(prompt.split("<engagement>\n")[1].split("\n</engagement>")[0])
    first = context["most_severe_findings"][0]
    assert first["detail"] == "Inventory count not attended" and first["account"] == "1300"
    assert "unclassified_accounts" in context["uncorrected_misstatements"]


def test_drafts_go_with_their_engagement(fake_api):  # noqa: F811
    eid = engagement()
    fake_api["answers"].append(RISKS)
    d = client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={}).json()
    assert client.delete(f"/api/engagements/{eid}").status_code == 204
    assert client.get(f"/api/ai-drafts/{d['id']}").status_code == 404


def test_busy_and_failed_are_logged(fake_api, monkeypatch):  # noqa: F811
    eid = engagement()
    fake_api["answers"].append(assistant.AssistantError("The Claude API is unavailable."))
    assert client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={}).status_code == 502
    monkeypatch.setattr(assistant, "_slots", assistant.threading.BoundedSemaphore(1))
    assistant._slots.acquire()
    try:
        r = client.post(f"/api/engagements/{eid}/ai/risk-assessment", json={})
    finally:
        assistant._slots.release()
    assert r.status_code == 429
    summaries = [e["summary"] for e in activity(eid, "assistant.risk_assessment")]
    assert summaries == ["Asked the AI assistant for a risk assessment (busy)",
                         "Asked the AI assistant for a risk assessment (failed)"]


@pytest.mark.parametrize("text", [
    "Asked the AI assistant to draft a management letter from 12 findings",
    "Asked the AI assistant to draft a management letter from 12 findings (busy)",
    "Asked the AI assistant to draft a management letter from 12 findings (failed)",
    "Asked the AI assistant for a risk assessment (failed)",
    "Downloaded AI risk assessment #3", "Downloaded AI management letter #4",
    "Written from 80 of 120 findings, the most severe first. The rest were not sent.",
    "There are no findings with those statuses to write about.",
    "The AI's answer had no risks in it. Try again.",
    "The AI's answer had no letter points in it. Try again.",
    *ai_writer.LEVELS.values(), *ai_writer.ASSERTIONS.values(), *ai_writer.AREAS.values(),
])
def test_new_text_has_arabic(text):
    from tests.test_ai_tools import translated  # noqa: PLC0415
    assert translated(text)
