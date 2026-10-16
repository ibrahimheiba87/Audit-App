"""The audit file pack, firm templates for the check settings, and run summaries."""

import io
import zipfile

import pytest
from openpyxl import load_workbook

from app import db
from tests.test_match_disclosures import EMPLOYEES, VENDORS, upload
from tests.test_statements import client, engagement


def actions():
    with db.SessionLocal() as s:
        return [a.action for a in s.query(db.Activity).order_by(db.Activity.id)]


def run_checks(eid, **settings):
    gl = upload(eid, "gl.csv", open("sample_data/sample_gl.csv", "rb").read(), kind="gl")
    r = client.post(f"/api/engagements/{eid}/runs", json={"upload_id": gl, **settings})
    assert r.status_code == 201, r.text
    return r.json()


def test_the_pack_holds_every_workpaper_and_an_index():
    eid = engagement(client_name="Pack Co")
    run_checks(eid)
    vendors, staff = upload(eid, "vendors.csv", VENDORS), upload(eid, "staff.csv", EMPLOYEES)
    client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": staff,
        "keys": [{"left": "Bank Account", "right": "IBAN", "method": "exact"}]})
    client.put(f"/api/engagements/{eid}/disclosures",
               json={"answers": {"ias1_compliance": {"status": "yes"}}})
    r = client.get(f"/api/engagements/{eid}/pack")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    index = [n for n in names if n.endswith("-index.xlsx")]
    assert len(index) == 1 and "/" not in index[0]
    folders = {n.split("/")[0] for n in names if "/" in n}
    assert folders == {"2 Testing", "3 Completion"}
    assert any("audit-report" in n for n in names)
    assert any("match" in n for n in names)
    assert any("lead-sheets" in n for n in names) and any("disclosures" in n for n in names)
    assert not any(n.startswith("4 Source files") for n in names)
    book = load_workbook(io.BytesIO(z.read(index[0])))
    assert {"About", "Contents", "Sign-offs", "Review notes", "Files", "Activity"} <= \
        set(book.sheetnames)
    listed = {row[1] for row in book["Contents"].iter_rows(min_row=2, values_only=True)}
    assert {n.split("/", 1)[1] for n in names if "/" in n} <= listed
    files = [row[0] for row in book["Files"].iter_rows(min_row=2, values_only=True)]
    assert "vendors.csv" in files and "gl.csv" in files
    assert actions()[-1] == "pack.download"


def test_the_pack_can_include_the_source_files():
    eid = engagement(client_name="Files Co", with_prior=False)
    upload(eid, "vendors.csv", VENDORS)
    upload(eid, "copy/vendors.csv", VENDORS)  # the same name twice gets a number
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/engagements/{eid}/pack?files=true").content))
    sources = sorted(n for n in z.namelist() if n.startswith("4 Source files/"))
    assert "4 Source files/vendors.csv" in sources and len(sources) == 3  # and the tb
    assert z.read("4 Source files/vendors.csv") == VENDORS


def test_a_broken_workpaper_is_listed_not_fatal(monkeypatch):
    from app import crossmatch

    eid = engagement(client_name="Broken Co", with_prior=False)
    vendors, staff = upload(eid, "vendors.csv", VENDORS), upload(eid, "staff.csv", EMPLOYEES)
    client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": staff,
        "keys": [{"left": "Bank Account", "right": "IBAN", "method": "exact"}]})

    def broken(m):
        raise KeyError("old result")
    monkeypatch.setattr(crossmatch, "workbook", broken)
    r = client.get(f"/api/engagements/{eid}/pack")
    assert r.status_code == 200, r.text
    z = zipfile.ZipFile(io.BytesIO(r.content))
    index = next(n for n in z.namelist() if n.endswith("-index.xlsx"))
    book = load_workbook(io.BytesIO(z.read(index)))
    missing = [row[0] for row in book["Not included"].iter_rows(min_row=2, values_only=True)]
    assert any(m.startswith("File match") for m in missing)


def test_firm_templates():
    r = client.post("/api/test-templates", json={
        "name": "Gulf retail", "cutoff_days": 7, "round_threshold": 5000, "weekend": "fri-sat",
        "keywords": ["cash", "director", "Cash"], "is_default": True})
    assert r.status_code == 201, r.text
    t = r.json()
    assert t["keywords"] == ["cash", "director"] and t["is_default"]
    assert client.post("/api/test-templates", json={"name": "Gulf retail"}).status_code == 409
    other = client.post("/api/test-templates", json={"name": "Other", "is_default": True}).json()
    listed = {x["name"]: x for x in client.get("/api/test-templates").json()}
    assert not listed["Gulf retail"]["is_default"] and listed["Other"]["is_default"]

    # A run with the template's settings as they are records its name; changed ones don't.
    eid = engagement(client_name="Template Co", with_prior=False)
    settings = {k: t[k] for k in ("cutoff_days", "round_threshold", "weekend", "keywords")}
    run = run_checks(eid, **settings, template_id=t["id"])
    assert run["settings"]["template"] == "Gulf retail"
    run = run_checks(eid, **{**settings, "cutoff_days": 3}, template_id=t["id"])
    assert "template" not in run["settings"]
    report = client.get(f"/api/runs/{run['id'] - 1}/report")
    assert report.status_code == 200

    unchanged = client.put(f"/api/test-templates/{other['id']}", json={
        "name": "Other", "is_default": True}).json()
    assert unchanged["updated_at"] == other["updated_at"]
    assert client.delete(f"/api/test-templates/{other['id']}").status_code == 204
    assert [a for a in actions() if a.startswith("template.")] == [
        "template.create", "template.create", "template.delete"]


@pytest.mark.usefixtures("real_auth")
def test_only_partners_change_templates():
    from tests.test_auth import add_user, admin
    preparer = add_user(admin(), "preparer")
    assert preparer.get("/api/test-templates").status_code == 200
    assert preparer.post("/api/test-templates", json={"name": "Mine"}).status_code == 403


def test_runs_list_from_their_summaries():
    eid = engagement(client_name="Summary Co", with_prior=False)
    run = run_checks(eid)
    with db.SessionLocal() as s:
        summary = s.get(db.RunSummary, run["id"])
        assert summary.rows == run["rows"]
        s.delete(summary)  # as a run saved before summaries existed
        s.commit()
    e = client.get(f"/api/engagements/{eid}").json()
    assert e["runs"][0]["rows"] == run["rows"]  # read from the result instead
    db.init()  # fills in the missing summary
    with db.SessionLocal() as s:
        assert s.get(db.RunSummary, run["id"]).findings_by_severity == \
            run["findings_by_severity"]
