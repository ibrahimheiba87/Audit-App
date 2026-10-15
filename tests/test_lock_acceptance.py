"""Locking and archiving engagements, and the acceptance and independence checklist."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import acceptance, db
from tests.test_auth import add_user, admin, join, setup_engagement
from tests.test_portal import PDF, add_client_user, ask, upload
from tests.test_review import section

pytestmark = pytest.mark.usefixtures("real_auth")


def actions(eid):
    with db.SessionLocal() as s:
        return [a.action for a in s.query(db.Activity).filter_by(engagement_id=eid)]


def crew():
    """Admin, preparer, reviewer and partner on one engagement (with a ledger upload)."""
    a = admin()
    prep, rev, partner = (add_user(a, r) for r in ("preparer", "reviewer", "partner"))
    eid, uid = setup_engagement(a)
    join(a, eid, prep, rev, partner)
    return a, prep, rev, partner, eid, uid


# --- Locking and archiving -------------------------------------------------------


def test_a_locked_file_can_be_read_but_not_changed():
    a, prep, rev, partner, eid, uid = crew()
    assert prep.post(f"/api/engagements/{eid}/lock", json={}).status_code == 403
    assert rev.post(f"/api/engagements/{eid}/lock", json={}).status_code == 403
    r = partner.post(f"/api/engagements/{eid}/lock", json={"note": " Report signed 2027-03-01 "})
    assert r.status_code == 200, r.text
    lock = r.json()["lock"]
    assert lock["locked_by"] == "Partner" and lock["note"] == "Report signed 2027-03-01"
    assert partner.post(f"/api/engagements/{eid}/lock", json={}).status_code == 409

    # Changes are refused with a 409 that says why...
    for r in (prep.post(f"/api/engagements/{eid}/runs", json={"upload_id": uid}),
              prep.post(f"/api/engagements/{eid}/risks",
                        json={"area": "revenue", "title": "Cut-off", "level": "lower"}),
              rev.patch(f"/api/engagements/{eid}", json={"name": "Renamed"}),
              prep.delete(f"/api/uploads/{uid}"),
              prep.put(f"/api/engagements/{eid}/acceptance", json={}),
              prep.post(f"/api/engagements/{eid}/independence", json={"independent": True})):
        assert r.status_code == 409, r.text
        assert "locked" in r.json()["detail"]
    # ...while reading, downloading, analyses that change nothing, the team and rolling
    # forward still work.
    assert prep.get(f"/api/engagements/{eid}").json()["lock"]["locked_by"] == "Partner"
    assert prep.get(f"/api/engagements/{eid}/acceptance/export").status_code == 200
    assert prep.post(f"/api/uploads/{uid}/stratify", json={}).status_code == 200
    assert a.put(f"/api/engagements/{eid}/team", json={"user_ids": []}).status_code == 200
    join(a, eid, prep, rev, partner)
    r = rev.post(f"/api/engagements/{eid}/roll-forward", json={"name": "FY2027"})
    assert r.status_code == 201, r.text
    assert r.json()["lock"] is None  # next year's file starts unlocked

    # A locked engagement (and its client) can't be deleted.
    r = a.delete(f"/api/engagements/{eid}")
    assert r.status_code == 409 and "locked" in r.json()["detail"]
    client_id = a.get(f"/api/engagements/{eid}").json()["client"]["id"]
    assert a.delete(f"/api/clients/{client_id}").status_code == 409

    # Only an admin unlocks, giving a reason, and both are logged.
    assert partner.request("DELETE", f"/api/engagements/{eid}/lock",
                           json={"reason": "Typo"}).status_code == 403
    assert a.request("DELETE", f"/api/engagements/{eid}/lock",
                     json={"reason": "  "}).status_code == 422
    r = a.request("DELETE", f"/api/engagements/{eid}/lock", json={"reason": "Report reissued"})
    assert r.status_code == 200 and r.json() == {"lock": None}
    assert a.request("DELETE", f"/api/engagements/{eid}/lock",
                     json={"reason": "Again"}).status_code == 409
    assert prep.post(f"/api/engagements/{eid}/runs", json={"upload_id": uid}).status_code == 201
    assert {"engagement.lock", "engagement.unlock"} <= set(actions(eid))


def test_outsiders_still_get_not_found():
    a, prep, _, partner, eid, uid = crew()
    outsider = add_user(a, "partner", "outsider@example.com")
    partner.post(f"/api/engagements/{eid}/lock", json={})
    assert outsider.post(f"/api/engagements/{eid}/runs",
                         json={"upload_id": uid}).status_code == 404
    assert outsider.delete(f"/api/uploads/{uid}").status_code == 404
    assert outsider.post(f"/api/engagements/{eid}/lock", json={}).status_code == 404


def test_archive_needs_a_lock_and_hides_the_engagement_from_the_list():
    a, _, _, partner, eid, _ = crew()
    r = partner.post(f"/api/engagements/{eid}/archive", json={"archived": True})
    assert r.status_code == 409 and "Lock" in r.json()["detail"]
    partner.post(f"/api/engagements/{eid}/lock", json={})
    r = partner.post(f"/api/engagements/{eid}/archive", json={"archived": True})
    assert r.status_code == 200 and r.json()["lock"]["archived_by"] == "Partner"
    eng = a.get("/api/clients").json()[0]["engagements"][0]
    assert eng["locked"] and eng["archived"]
    r = partner.post(f"/api/engagements/{eid}/archive", json={"archived": False})
    assert r.json()["lock"]["archived"] is False and r.json()["lock"]["archived_at"] is None
    assert {"engagement.archive", "engagement.unarchive"} <= set(actions(eid))
    # Unlocking an archived engagement takes it out of the archive too.
    partner.post(f"/api/engagements/{eid}/archive", json={"archived": True})
    a.request("DELETE", f"/api/engagements/{eid}/lock", json={"reason": "Needed again"})
    eng = a.get("/api/clients").json()[0]["engagements"][0]
    assert not eng["locked"] and not eng["archived"]


def test_the_client_portal_is_locked_too():
    a = admin()
    prep = add_user(a, "preparer")
    partner = add_user(a, "partner")
    acme = a.post("/api/clients", json={"name": "Acme"}).json()
    eid = a.post(f"/api/clients/{acme['id']}/engagements",
                 json={"name": "FY2026", "period_end": "2026-12-31"}).json()["id"]
    cl = add_client_user(a, acme["id"])
    join(a, eid, prep, partner)
    rq = ask(prep, eid)
    partner.post(f"/api/engagements/{eid}/lock", json={})
    r = upload(cl, rq["id"], PDF)
    assert r.status_code == 409 and "locked" in r.json()["detail"]
    assert cl.get("/api/portal/requests").status_code == 200


def test_the_dashboard_has_no_next_steps_once_locked():
    _, prep, _, partner, eid, _ = crew()
    assert prep.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]
    partner.post(f"/api/engagements/{eid}/lock", json={})
    assert prep.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"] == []


# --- Acceptance and independence -------------------------------------------------


def all_answers(**changed):
    return {q: {"answer": changed.get(q, "yes" if expected else "no"), "comment": ""}
            for q, _, expected in acceptance.QUESTIONS}


def steps(c, eid):
    return [s["text"] for s in c.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]]


def test_checklist_decision_and_flags():
    _, prep, rev, partner, eid, _ = crew()
    empty = prep.get(f"/api/engagements/{eid}/acceptance").json()
    assert empty["answered"] == 0 and empty["decision"] is None
    assert len(empty["questions"]) == len(acceptance.QUESTIONS)
    assert not section(prep, eid, "acceptance")["needed"]  # optional until started

    r = prep.put(f"/api/engagements/{eid}/acceptance",
                 json={"answers": {"integrity": {"answer": "yes"}}})
    assert r.status_code == 200 and r.json()["answered"] == 1
    assert "Record the acceptance decision." in steps(prep, eid)
    assert section(prep, eid, "acceptance")["needed"]
    # Answers are merged, so one person's save doesn't wipe another's.
    r = rev.put(f"/api/engagements/{eid}/acceptance",
                json={"answers": {"fees": {"answer": "no", "comment": "Unpaid 2025 fee"}}})
    assert r.json()["answers"]["integrity"]["answer"] == "yes"
    assert r.json()["flags"] == ["fees"]
    assert prep.put(f"/api/engagements/{eid}/acceptance",
                    json={"answers": {"nonsense": {"answer": "yes"}}}).status_code == 422

    # Only a partner decides, once every question is answered...
    assert rev.post(f"/api/engagements/{eid}/acceptance/decision",
                    json={"decision": "accept"}).status_code == 403
    r = partner.post(f"/api/engagements/{eid}/acceptance/decision", json={"decision": "accept"})
    assert r.status_code == 400 and "left" in r.json()["detail"]
    prep.put(f"/api/engagements/{eid}/acceptance", json={"answers": all_answers(fees="no")})
    # ...and accepting with a flagged answer needs a reason.
    r = partner.post(f"/api/engagements/{eid}/acceptance/decision", json={"decision": "accept"})
    assert r.status_code == 400 and "second look" in r.json()["detail"]
    r = partner.post(f"/api/engagements/{eid}/acceptance/decision",
                     json={"decision": "accept", "note": "Fee paid after the year end."})
    assert r.status_code == 200, r.text
    s = r.json()
    assert s["decision"] == "accept" and s["decision_label"] == "Accepted"
    assert s["decided_by"] == "Partner" and not s["changed_since_decision"]
    assert "Record the acceptance decision." not in steps(prep, eid)
    prep.put(f"/api/engagements/{eid}/acceptance", json={"conclusion": "Fine."})
    assert prep.get(f"/api/engagements/{eid}/acceptance").json()["changed_since_decision"]

    # Declining is always allowed, and shows on the dashboard.
    partner.post(f"/api/engagements/{eid}/acceptance/decision", json={"decision": "decline"})
    assert "The engagement was declined at acceptance." in steps(prep, eid)
    assert {"acceptance.update", "acceptance.decide"} <= set(actions(eid))


def test_independence_declarations():
    _, prep, rev, partner, eid, _ = crew()
    r = prep.post(f"/api/engagements/{eid}/independence", json={"independent": False})
    assert r.status_code == 400 and "threat" in r.json()["detail"]
    r = prep.post(f"/api/engagements/{eid}/independence", json={"independent": True})
    assert r.status_code == 200
    ind = r.json()["independence"]
    assert ind["declared"] == 1 and ind["expected"] == 3 and ind["mine"]["independent"]
    assert "2 team members still to confirm independence." in steps(prep, eid)

    r = rev.post(f"/api/engagements/{eid}/independence", json={
        "independent": False, "threats": "My brother is the CFO.",
        "safeguards": "Not on revenue or payroll."})
    ind = r.json()["independence"]
    assert ind["threats"] == 1 and ind["mine"]["threats"] == "My brother is the CFO."
    assert "Consider the threats to independence declared by 1 team member." in steps(prep, eid)
    # Declaring again replaces the earlier declaration.
    r = rev.post(f"/api/engagements/{eid}/independence",
                 json={"independent": True, "threats": "ignored"})
    ind = r.json()["independence"]
    assert ind["threats"] == 0 and ind["mine"]["threats"] == ""
    partner.post(f"/api/engagements/{eid}/independence", json={"independent": True})
    assert not any("independence" in s for s in steps(prep, eid))
    assert "independence.declare" in actions(eid)


def test_acceptance_export():
    _, prep, _, partner, eid, _ = crew()
    prep.put(f"/api/engagements/{eid}/acceptance",
             json={"answers": all_answers(), "conclusion": "=SUM(A1)"})
    partner.post(f"/api/engagements/{eid}/acceptance/decision", json={"decision": "accept"})
    prep.post(f"/api/engagements/{eid}/independence", json={"independent": True})
    xl = prep.get(f"/api/engagements/{eid}/acceptance/export")
    assert xl.status_code == 200
    sheets = pd.read_excel(io.BytesIO(xl.content), sheet_name=None)
    assert list(sheets) == ["About", "Acceptance checklist", "Independence"]
    assert len(sheets["Acceptance checklist"]) == len(acceptance.QUESTIONS)
    book = load_workbook(io.BytesIO(xl.content))  # the conclusion stays text, not a formula
    assert not [c for ws in book for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert any(c.value == "=SUM(A1)" for ws in book for row in ws.iter_rows() for c in row)
    assert "acceptance.download" in actions(eid)
