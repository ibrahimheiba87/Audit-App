"""Matching two files, and the IFRS disclosure checklist with the AI reading the statements."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import crossmatch, db, disclosures
from tests.test_ai_tools import fake_api  # noqa: F401 (fixture)
from tests.test_statements import client, engagement

VENDORS = b"""Vendor No,Vendor Name,Bank Account,Phone,Address
V1,Gulf Supplies LLC,AE07 0331 2345 6789 0123 456,+971 4 555 0101,12 Palm Street
V2,Desert Trading Co,AE11 0222 0000 1111 2222 333,04-555-0199,8 Creek Road
V3,Ahmed Al Mansouri,AE55 0123 9999 8888 7777 666,050 111 2222,Villa 3 Jumeirah
V4,Gulf Suplies,AE07033123456789012345 6,,
V5,Northwind,,,
"""
EMPLOYEES = b"""Staff ID,Name,IBAN,Mobile,Home address
E1,Ahmed Al-Mansouri,AE55-0123-9999-8888-7777-666,+971 50 111 2222,"villa 3, jumeirah"
E2,Sara Khan,AE90 0444 1234 1234 1234 123,055 999 0000,Flat 4 Marina
"""


def upload(eid, name, content, kind="list"):
    r = client.post(f"/api/engagements/{eid}/uploads", data={"kind": kind},
                    files={"file": (name, content, "text/csv")})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def actions(eid):
    with db.SessionLocal() as s:
        return [a.action for a in s.query(db.Activity).filter_by(engagement_id=eid)]


# --- Matching two files ------------------------------------------------------------


def test_normalizing():
    assert crossmatch.normalize("AE07 0331-2345", "exact") == "ae0703312345"
    assert crossmatch.normalize("+971 (4) 555-0101", "digits") == "97145550101"
    assert crossmatch.normalize("12", "digits") == ""  # too little to match on
    assert crossmatch.normalize("The Gulf Supplies Co. LLC", "name") == "gulf supplies"
    assert crossmatch.normalize(1234567.0, "exact") == "1234567"
    assert crossmatch.normalize(None, "exact") == ""


def test_vendors_against_employees():
    eid = engagement(with_prior=False)
    vendors, staff = upload(eid, "vendors.csv", VENDORS), upload(eid, "staff.csv", EMPLOYEES)
    cols = client.get(f"/api/uploads/{vendors}/columns").json()
    assert cols["columns"] == ["Vendor No", "Vendor Name", "Bank Account", "Phone", "Address"]
    assert cols["rows"] == 5
    r = client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": staff, "title": "Vendors v staff",
        "keys": [{"left": "Bank Account", "right": "IBAN", "method": "exact"},
                 {"left": "Vendor Name", "right": "Name", "method": "name"},
                 {"left": "Phone", "right": "Mobile", "method": "digits"},
                 {"left": "Address", "right": "Home address", "method": "exact"}]})
    assert r.status_code == 201, r.text
    m = r.json()
    res = m["result"]
    assert res["totals"]["pairs"] == 1 and res["totals"]["all_keys"] == 0
    [x] = res["matches"]
    # Ahmed: same IBAN and address however they're written, and a similar name; the phone
    # numbers differ by the country code, so they don't match as written.
    assert (x["left_row"], x["right_row"]) == (4, 2)
    assert [k["key"] for k in x["keys"]] == [0, 1, 3]
    assert x["left"]["Vendor Name"] == "Ahmed Al Mansouri"
    assert [k["pairs"] for k in res["keys"]] == [1, 1, 0, 1]
    eng = client.get(f"/api/engagements/{eid}").json()
    assert eng["cross_matches"][0]["title"] == "Vendors v staff"

    xl = client.get(f"/api/cross-matches/{m['id']}/export")
    sheets = pd.read_excel(io.BytesIO(xl.content), sheet_name=None)
    assert list(sheets) == ["About", "Columns", "Matches"]
    assert sheets["Matches"]["1: Vendor No"].tolist() == ["V3"]
    assert client.delete(f"/api/cross-matches/{m['id']}").status_code == 204
    assert {"crossmatch.create", "crossmatch.download", "crossmatch.delete"} <= set(actions(eid))


def test_duplicates_in_one_file():
    eid = engagement(with_prior=False)
    vendors = upload(eid, "vendors.csv", VENDORS)
    r = client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": vendors, "threshold": 85,
        "keys": [{"left": "Bank Account", "right": "Bank Account", "method": "exact"},
                 {"left": "Vendor Name", "right": "Vendor Name", "method": "name"}]})
    res = r.json()["result"]
    # V1 and V4: the same account and a misspelt name. No row matches itself.
    assert [(x["left_row"], x["right_row"]) for x in res["matches"]] == [(2, 5)]
    assert res["same_file"] and len(res["matches"][0]["keys"]) == 2


def test_matching_is_checked():
    eid = engagement(with_prior=False)
    vendors = upload(eid, "vendors.csv", VENDORS)
    key = {"left": "Nope", "right": "Vendor Name", "method": "exact"}
    r = client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": vendors, "keys": [key]})
    assert r.status_code == 400 and "Nope" in r.json()["detail"]
    for bad in ({"keys": []}, {"threshold": 50}, {"keys": [{**key, "method": "fuzzy"}]}):
        body = {"left_upload_id": vendors, "right_upload_id": vendors, "keys": [key], **bad}
        assert client.post(f"/api/engagements/{eid}/cross-matches",
                           json=body).status_code == 422
    other = engagement(with_prior=False, client_name="Other")
    theirs = upload(other, "v.csv", VENDORS)
    r = client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": theirs,
        "keys": [{"left": "Vendor No", "right": "Vendor No"}]})
    assert r.status_code == 400


def test_lists_are_not_ledgers():
    eid = engagement(with_prior=False)
    vendors = upload(eid, "vendors.csv", VENDORS)
    r = client.post(f"/api/engagements/{eid}/runs", json={"upload_id": vendors})
    assert r.status_code == 400 and "list with no amounts" in r.json()["detail"]
    assert client.post(f"/api/uploads/{vendors}/stratify", json={}).status_code == 400
    up = next(u for u in client.get(f"/api/engagements/{eid}").json()["uploads"]
              if u["id"] == vendors)
    assert up["import"]["rows_used"] == 5 and up["mapping"] is None
    r = client.post(f"/api/engagements/{eid}/uploads", data={"kind": "list"},
                    files={"file": ("empty.csv", b"A,B\n", "text/csv")})
    assert r.status_code == 400


def test_deleting_a_file_deletes_its_matches():
    eid = engagement(with_prior=False)
    vendors, staff = upload(eid, "vendors.csv", VENDORS), upload(eid, "staff.csv", EMPLOYEES)
    client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": vendors, "right_upload_id": staff,
        "keys": [{"left": "Bank Account", "right": "IBAN"}]})
    assert client.delete(f"/api/uploads/{staff}").status_code == 204
    assert client.get(f"/api/engagements/{eid}").json()["cross_matches"] == []


def test_many_similar_names_stay_quick():
    left = pd.DataFrame({"n": [f"Supplier {i} Trading" for i in range(100, 3100)]})
    right = pd.DataFrame({"n": [f"Suplier {i} Trading" for i in range(100, 3100, 7)]})
    res = crossmatch.match(left, right, [{"left": "n", "right": "n", "method": "name"}], 90)
    assert res["totals"]["pairs"] >= len(right)


# --- Disclosure checklist ------------------------------------------------------------

PDF = ("fs.pdf", b"%PDF-1.4 draft statements", "application/pdf")


def test_checklist_answers_and_dashboard():
    eid = engagement(with_prior=False)
    s = client.get(f"/api/engagements/{eid}/disclosures").json()
    assert s["counts"] == {"items": len(disclosures.ITEMS), "answered": 0, "given": 0,
                           "missing": 0, "na": 0}
    steps = [x["text"] for x in client.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]]
    assert not any("disclosure" in x for x in steps)  # optional until started
    r = client.put(f"/api/engagements/{eid}/disclosures", json={"answers": {
        "ias1_compliance": {"status": "yes", "reference": "Note 2"},
        "ias24_related": {"status": "no", "comment": "No KMP note"}}})
    assert r.status_code == 200, r.text
    assert r.json()["counts"]["answered"] == 2
    steps = [x["text"] for x in client.get(f"/api/engagements/{eid}/dashboard").json()["next_steps"]]
    assert "Have the client add 1 missing disclosure to the statements." in steps
    assert f"Finish the disclosure checklist ({len(disclosures.ITEMS) - 2} left)." in steps
    # Saving the same again changes nothing and isn't logged twice.
    client.put(f"/api/engagements/{eid}/disclosures", json={"answers": {
        "ias1_compliance": {"status": "yes", "reference": "Note 2"}}})
    assert actions(eid).count("disclosure.update") == 1
    r = client.put(f"/api/engagements/{eid}/disclosures", json={"answers": {"x": {}}})
    assert r.status_code == 400 and r.json()["detail"] == "Unknown checklist item: x."


def test_ai_reads_the_statements_and_suggestions_are_accepted(fake_api):  # noqa: F811
    eid = engagement(with_prior=False)
    client.put(f"/api/engagements/{eid}/disclosures", json={"answers": {
        "ias1_compliance": {"status": "no"}}})
    fake_api["answers"].append({"items": [
        {"id": "ias1_compliance", "status": "present", "where": "Note 2",
         "evidence": "These statements comply with IFRS."},
        {"id": "ifrs16_leases", "status": "not_applicable", "where": "", "evidence": "No leases."},
        {"id": "ias24_related", "status": "missing", "where": "", "evidence": "Not found."},
        {"id": "made_up", "status": "present"}, {"id": "ias7_cash", "status": "maybe"}, "junk"]})
    r = client.post(f"/api/engagements/{eid}/disclosures/ai-review", files={"file": PDF},
                    data={"language": "ar"})
    assert r.status_code == 201, r.text
    s = r.json()
    system, content, _ = fake_api["calls"][0]
    assert system == disclosures.SYSTEM and content[0]["type"] == "document"
    assert "ias24_related (IAS 24)" in content[1]["text"] and "Arabic" in content[1]["text"]
    items = {x["id"]: x for x in s["items"]}
    assert items["ias1_compliance"]["differs"]  # the team said missing; the AI found it
    assert items["ias7_cash"]["ai"] is None
    assert s["ai"]["filename"] == "fs.pdf" and s["ai_pending"] == 2 and s["differs"] == 1
    assert s["counts"]["answered"] == 1  # suggestions aren't answers

    s = client.post(f"/api/engagements/{eid}/disclosures/accept-ai").json()
    items = {x["id"]: x for x in s["items"]}
    assert items["ias1_compliance"]["status"] == "no"  # answers given stay
    assert items["ifrs16_leases"]["status"] == "na"
    assert items["ias24_related"]["status"] == "no"
    assert items["ias24_related"]["comment"] == "AI: Not found."
    assert {"disclosure.ai_review", "disclosure.accept_ai"} <= set(actions(eid))

    xl = client.get(f"/api/engagements/{eid}/disclosures/export")
    book = load_workbook(io.BytesIO(xl.content))
    assert book.sheetnames == ["About", "Disclosure checklist"]
    assert book["Disclosure checklist"].max_row == len(disclosures.ITEMS) + 1


@pytest.mark.parametrize("content", [b"not a pdf", b"\x89PNG image"])
def test_ai_review_needs_a_pdf(fake_api, content):  # noqa: F811
    eid = engagement(with_prior=False)
    r = client.post(f"/api/engagements/{eid}/disclosures/ai-review",
                    files={"file": ("x.pdf", content, "application/pdf")})
    assert r.status_code == 400
    assert not fake_api["calls"]
    assert client.post(f"/api/engagements/{eid}/disclosures/accept-ai").status_code == 400


# --- Review fixes ------------------------------------------------------------------


def test_names_with_dotted_suffixes():
    assert crossmatch.normalize("Gulf Supplies L.L.C.", "name") == "gulf supplies"
    assert crossmatch.normalize("Acme Trading W.L.L.", "name") == "acme"
    assert crossmatch.normalize("J. Smith", "name") == "smith"


def test_common_values_are_skipped_not_paired():
    """A value in every row (a city, "N/A") would pair every row with every other."""
    rows = 3000
    raw = pd.DataFrame({"City": ["Riyadh"] * rows, "Phone": [f"0500{i:06d}" for i in range(rows)]})
    raw.loc[5, "Phone"] = raw.loc[9, "Phone"]
    r = crossmatch.match(raw, raw, [{"left": "City", "right": "City", "method": "exact"},
                                    {"left": "Phone", "right": "Phone", "method": "digits"}],
                         90, same_file=True)
    assert r["totals"]["pairs"] == 1
    assert r["keys"][0]["common_values"] == [{"value": "riyadh", "rows": rows}]
    assert r["totals"]["left_matched"] == r["totals"]["right_matched"] == 2  # rows 5 and 9


def test_too_many_pairs_is_refused():
    raw = pd.DataFrame({"Code": [f"A{i % 1500:04d}" for i in range(60_000)]})
    with pytest.raises(crossmatch.MatchError, match="tell the rows apart"):
        crossmatch.match(raw, raw, [{"left": "Code", "right": "Code", "method": "exact"}], 90,
                         same_file=True)


def test_digits_keep_leading_zeros_and_key_columns_beyond_the_first_30():
    eid = engagement(client_name="Zeros", with_prior=False)
    wide = ",".join(f"C{i}" for i in range(35))
    left = upload(eid, "a.csv", (wide + ",Account\n" + ",".join("x" * 35) + ",0012345678\n").encode())
    right = upload(eid, "b.xlsx", _xlsx({"Account": ["0012345678"]}))
    r = client.post(f"/api/engagements/{eid}/cross-matches", json={
        "left_upload_id": left, "right_upload_id": right,
        "keys": [{"left": "Account", "right": "Account", "method": "exact"}]})
    assert r.status_code == 201, r.text
    m = r.json()
    assert m["totals"]["pairs"] == 1
    assert m["result"]["matches"][0]["left"]["Account"] == "0012345678"
    assert "Account" in m["result"]["left_columns"]


def _xlsx(columns: dict) -> bytes:
    buf = io.BytesIO()
    pd.DataFrame(columns).astype(str).to_excel(buf, index=False)
    return buf.getvalue()


def test_list_on_a_second_sheet():
    eid = engagement(client_name="Sheets", with_prior=False)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xl:
        pd.DataFrame({"Cover": ["Vendor master"]}).to_excel(xl, sheet_name="Cover", index=False)
        pd.DataFrame({"Vendor": ["A", "B"]}).to_excel(xl, sheet_name="Vendors", index=False)
    r = client.post(f"/api/engagements/{eid}/uploads", data={"kind": "list", "sheet": "Vendors"},
                    files={"file": ("v.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    cols = client.get(f"/api/uploads/{r.json()['id']}/columns").json()
    assert cols["columns"] == ["Vendor"] and cols["rows"] == 2


def test_one_file_with_different_columns_keeps_direction():
    """Row 2's phone is row 1's alternate phone: the pair is shown that way round."""
    raw = pd.DataFrame({"Name": ["A", "B", "C"], "Phone": ["11111111", "22222222", "33333333"],
                        "Alternate": ["22222222", "99999999", "11111111"]})
    keys = [{"left": "Phone", "right": "Alternate", "method": "digits"}]
    r = crossmatch.match(raw, raw, keys, 90, same_file=True)
    got = {(m["left"]["Name"], m["right"]["Name"]) for m in r["matches"]}
    assert got == {("B", "A"), ("A", "C")}
    for m in r["matches"]:
        assert m["left"]["Phone"] == m["right"]["Alternate"]
    # A column matched against itself joins the pair whichever way round it was found.
    raw["City"] = ["Dubai", "Dubai", "Abu Dhabi"]
    keys.append({"left": "City", "right": "City", "method": "exact"})
    r = crossmatch.match(raw, raw, keys, 90, same_file=True)
    both = [m for m in r["matches"] if len(m["keys"]) == 2]
    assert [(m["left"]["Name"], m["right"]["Name"]) for m in both] == [("B", "A")]
    assert r["totals"]["all_keys"] == 1 and r["totals"]["pairs"] == 2
    # Two rows that each hold the other's phone match both ways, and both pairs get the city.
    raw = pd.DataFrame({"Name": ["A", "B"], "Phone": ["11111111", "22222222"],
                        "Alternate": ["22222222", "11111111"], "City": ["Dubai", "Dubai"]})
    r = crossmatch.match(raw, raw, keys, 90, same_file=True)
    assert r["totals"]["pairs"] == 2 and r["totals"]["all_keys"] == 2
